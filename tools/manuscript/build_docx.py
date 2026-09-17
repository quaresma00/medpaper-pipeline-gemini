#!/usr/bin/env python3
"""Build and audit journal-ready DOCX files with deterministic medical-manuscript styling."""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import tomllib
import zipfile
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - clear deployment failure
    raise SystemExit("python-docx is required; run bootstrap.py to install pinned dependencies") from exc


ROOT = Path(__file__).resolve().parents[2]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "rel": PKG_REL}
TEXT_ROLES = {"manuscript", "title_page", "cover_letter", "supplementary",
              "statements", "figure_legends"}
FORBIDDEN_PARAGRAPH_CONTROLS = (
    "outlineLvl", "keepNext", "keepLines", "pageBreakBefore",
)
THEMATIC_BREAK_RE = re.compile(
    r"(?m)^\s{0,3}(?:(?:\*\s*){3,}|(?:-\s*){3,}|(?:_\s*){3,})\s*$"
)
PIPE_TABLE_DIVIDER_RE = re.compile(
    r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$"
)


def project_root() -> Path:
    raw = os.environ.get("MEDPAPER_PROJECT")
    return Path(raw).resolve() if raw else ROOT / "project"


def planned_figure_ids(project: Path | None = None) -> list[str] | None:
    """Return every planned main/supplementary figure id, or None if the plan is unusable."""
    path = (project or project_root()) / "01_protocol/artifact_plan.json"
    if not path.is_file():
        return None
    try:
        plan = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    return [str(item.get("id", "")).strip()
            for group in ("main_figures", "supp_figures")
            for item in plan.get(group, []) if str(item.get("id", "")).strip()]


def figure_legend_cap() -> int:
    try:
        data = tomllib.loads((ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8"))
        return int(data.get("targets", {}).get("figure_legend_words_max", 180))
    except (OSError, ValueError, tomllib.TOMLDecodeError):
        return 180


def load_style(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    defaults = {
        "font_family": "Times New Roman", "body_font_pt": 12,
        "title_font_pt": 14, "section_heading_font_pt": 12,
        "subsection_heading_font_pt": 12, "line_spacing": 2.0,
        "paragraph_spacing_after_pt": 0, "margins_in": 1.0,
        "paper_size": "A4", "all_text_black": True,
        "external_hyperlinks": False,
    }
    defaults.update(data)
    return defaults


def _ensure_style(doc, name: str, base: str, font: str, size: float,
                  bold: bool = False, italic: bool = False):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    ppr = style.element.get_or_add_pPr()
    for tag in (*[f"w:{name}" for name in FORBIDDEN_PARAGRAPH_CONTROLS], "w:numPr"):
        node = ppr.find(qn(tag))
        if node is not None:
            ppr.remove(node)
    return style


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def iter_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _flatten_hyperlinks(doc) -> None:
    for hyperlink in list(doc.element.xpath(".//w:hyperlink")):
        parent = hyperlink.getparent()
        idx = parent.index(hyperlink)
        for child in list(hyperlink):
            hyperlink.remove(child)
            parent.insert(idx, child)
            idx += 1
        parent.remove(hyperlink)
    for part in doc.part.package.parts:
        for rel_id, rel in list(part.rels.items()):
            if rel.reltype == RELATIONSHIP_TYPE.HYPERLINK:
                part.drop_rel(rel_id)


def _replace_manual_line_breaks(doc) -> None:
    """Convert manual line breaks to real paragraphs without merging address/list lines."""
    for part in doc.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            continue
        for element in list(root.iter(qn("w:br"))) + list(root.iter(qn("w:cr"))):
            if element.tag == qn("w:br") and element.get(qn("w:type")) not in (None, "textWrapping"):
                continue
            paragraph = next((node for node in element.iterancestors() if node.tag == qn("w:p")), None)
            if paragraph is None:
                raise ValueError("manual line break outside a paragraph cannot be converted safely")
            current, tail = element, None
            while current is not paragraph:
                parent = current.getparent()
                following = list(parent)[parent.index(current) + 1:]
                clone = deepcopy(parent)
                for child in list(clone):
                    if child.tag not in {qn("w:pPr"), qn("w:rPr")}:
                        clone.remove(child)
                if tail is not None:
                    clone.append(tail)
                for child in following:
                    clone.append(child)
                current, tail = parent, clone
            paragraph.addnext(tail)
            # A section break belongs to the final resulting paragraph only.
            ppr = paragraph.find(qn("w:pPr"))
            sect = None if ppr is None else ppr.find(qn("w:sectPr"))
            if sect is not None:
                ppr.remove(sect)
            element.getparent().remove(element)


def _remove_heading_metadata(paragraph, *, strip_numbering: bool = True) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    tags = [f"w:{name}" for name in FORBIDDEN_PARAGRAPH_CONTROLS]
    if strip_numbering:
        tags.append("w:numPr")
    for tag in tags:
        node = ppr.find(qn(tag))
        if node is not None:
            ppr.remove(node)


def _source_problems(kind: str, inputs: list[Path]) -> list[str]:
    text = "\n\n".join(path.read_text(encoding="utf-8", errors="replace") for path in inputs)
    problems: list[str] = []
    if kind in {"manuscript", "figure_legends"}:
        count = len(re.findall(r"(?mi)^#\s+Figure legends\s*$", text))
        if count != 1:
            problems.append(f"source must contain exactly one '# Figure legends' heading, found {count}")
    if kind == "supplementary":
        lines = text.splitlines()
        if (any(PIPE_TABLE_DIVIDER_RE.match(line) for line in lines) or
                re.search(r"<table\b", text, re.I | re.S) or
                re.search(r"^\s*\+[=-]{3,}(?:\+[=-]{3,})+\+\s*$", text, re.M)):
            problems.append(
                "supplementary Methods contains a table; move it to the planned supplementary "
                "tables workbook and build it with tools/tables/threeline.py"
            )
        if THEMATIC_BREAK_RE.search(text):
            problems.append("supplementary Methods contains a Markdown thematic break/horizontal rule")
    return problems


def _set_run(run, font: str, size: float) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    rstyle = rpr.find(qn("w:rStyle"))
    if rstyle is not None and rstyle.get(qn("w:val"), "").casefold() == "hyperlink":
        rpr.remove(rstyle)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    color = rpr.find(qn("w:color"))
    if color is not None:
        color.set(qn("w:val"), "000000")
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attr}"), None)


def normalize_docx(path: Path, style: dict, kind: str) -> None:
    doc = Document(path)
    font = str(style["font_family"])
    body_pt = float(style["body_font_pt"])
    title_pt = float(style["title_font_pt"])
    h1_pt = float(style["section_heading_font_pt"])
    h2_pt = float(style["subsection_heading_font_pt"])
    spacing = float(style["line_spacing"])
    after = float(style["paragraph_spacing_after_pt"])

    body_style = _ensure_style(doc, "Manuscript Body", "Normal", font, body_pt)
    title_style = _ensure_style(doc, "Manuscript Title", "Normal", font, title_pt, bold=True)
    # Both Markdown section levels are flattened onto one Normal-based paragraph
    # style.  Typography may still reflect the journal's H1/H2 sizes, but Word has
    # no outline hierarchy to display or collapse.
    section_style = _ensure_style(doc, "SectionHeading", "Normal", font, h1_pt, bold=True)
    for builtin in ("Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2",
                    "Heading 3", "Bibliography", "Hyperlink"):
        try:
            s = doc.styles[builtin]
        except KeyError:
            continue
        s.font.name = font
        s.font.color.rgb = RGBColor(0, 0, 0)
        if builtin == "Hyperlink":
            s.font.underline = False
        ppr = s.element.get_or_add_pPr()
        for tag in ("w:outlineLvl", "w:keepNext", "w:keepLines", "w:pageBreakBefore"):
            node = ppr.find(qn(tag))
            if node is not None:
                ppr.remove(node)

    _flatten_hyperlinks(doc)
    _replace_manual_line_breaks(doc)
    paras = list(iter_paragraphs(doc))
    nonempty_seen = 0
    for p in paras:
        original = p.style.name if p.style is not None else ""
        text = p.text.strip()
        if text:
            nonempty_seen += 1
        _remove_heading_metadata(p, strip_numbering=False)
        if original == "Title" or (kind == "manuscript" and nonempty_seen == 1):
            p.style = title_style
            size = title_pt
            _remove_heading_metadata(p)
        elif re.match(r"^Heading 1$", original, re.I):
            p.style = section_style
            size = h1_pt
            _remove_heading_metadata(p)
        elif re.match(r"^Heading [2-9]$", original, re.I):
            p.style = section_style
            size = h2_pt
            _remove_heading_metadata(p)
        else:
            if original in {"Normal", "Body Text", "First Paragraph", "Compact"}:
                p.style = body_style
            size = body_pt
        p.paragraph_format.line_spacing = spacing
        p.paragraph_format.space_after = Pt(after)
        for run in p.runs:
            _set_run(run, font, size)

    margin = Inches(float(style["margins_in"]))
    for section in doc.sections:
        section.top_margin = section.bottom_margin = margin
        section.left_margin = section.right_margin = margin
        if str(style["paper_size"]).upper() == "LETTER":
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        else:
            section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
        section.orientation = WD_ORIENT.PORTRAIT

    doc.save(path)
    _sanitize_package_xml(path, font)


def _sanitize_package_xml(path: Path, font: str) -> None:
    """Remove theme residue and every Word paragraph control the user prohibited."""
    with zipfile.ZipFile(path, "r") as src:
        entries = [(item, src.read(item.filename)) for item in src.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item, blob in entries:
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                # Word's mc:Ignorable attributes contain namespace-prefix strings; parsing
                # and reserializing them with ElementTree can rename prefixes and make Word
                # report a corrupt file. Byte replacement preserves the package exactly.
                blob = re.sub(br"Aptos(?: Display)?", font.encode("utf-8"), blob,
                              flags=re.I)
                for tag in (*FORBIDDEN_PARAGRAPH_CONTROLS, "pBdr"):
                    name = tag.encode("ascii")
                    blob = re.sub(rb"<w:" + name + rb"\b[^>]*/>", b"", blob)
                    blob = re.sub(rb"<w:" + name + rb"\b[^>]*>.*?</w:" + name + rb">",
                                  b"", blob, flags=re.S)
            out.writestr(item, blob)
    tmp.replace(path)


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    return "\n".join("".join(t.text or "" for t in p.findall(".//w:t", NS))
                       for p in root.findall(".//w:p", NS))


def _canon_figure_id(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().rstrip(".")).casefold()


def _legend_problems(root, expected: list[str], cap: int) -> list[str]:
    paragraphs: list[tuple[str, str]] = []
    for paragraph in root.findall(".//w:p", NS):
        paragraph_text = "".join(t.text or "" for t in paragraph.findall(".//w:t", NS)).strip()
        pstyle = paragraph.find("w:pPr/w:pStyle", NS)
        style_id = "" if pstyle is None else pstyle.get(f"{{{W}}}val", "")
        if paragraph_text:
            paragraphs.append((paragraph_text, style_id))
    try:
        start = next(i for i, (value, _) in enumerate(paragraphs)
                     if value.casefold() == "figure legends") + 1
    except StopIteration:
        return ["Figure legends heading is missing"]
    headings: list[tuple[int, str, str]] = []
    for i in range(start, len(paragraphs)):
        value, style_id = paragraphs[i]
        match = re.match(r"^(Figure\s+S?\d+)\.?\s*(.*)$", value, re.I)
        if match and style_id == "SectionHeading":
            headings.append((i, match.group(1), match.group(2)))
    blocks: dict[str, tuple[int, int]] = {}
    for j, (idx, figure_id, suffix) in enumerate(headings):
        end = headings[j + 1][0] if j + 1 < len(headings) else len(paragraphs)
        content = "\n".join([suffix] + [value for value, _ in paragraphs[idx + 1:end]]).strip()
        blocks[_canon_figure_id(figure_id)] = (
            len(content), len(re.findall(r"\b[A-Za-z][A-Za-z0-9'-]*\b", content)))
    wanted = {_canon_figure_id(value): value for value in expected}
    problems: list[str] = []
    seen = [_canon_figure_id(item[1]) for item in headings]
    duplicates = sorted({item for item in seen if seen.count(item) > 1})
    if duplicates:
        problems.append("duplicate figure legend heading(s): " + ", ".join(duplicates))
    for idx, figure_id, _ in headings:
        if idx + 1 < len(paragraphs):
            first_body = paragraphs[idx + 1][0]
            if re.match(rf"^{re.escape(figure_id)}\b", first_body, re.I):
                problems.append(f"{figure_id} is repeated at the start of its legend body")
    for key, display in wanted.items():
        if key not in blocks:
            problems.append(f"Figure legends lacks {display}")
            continue
        chars, words = blocks[key]
        if chars < 40:
            problems.append(f"{display} legend is only {chars} characters")
        if words > cap:
            problems.append(f"{display} legend is {words} words (maximum {cap})")
    extras = sorted(set(blocks) - set(wanted))
    if extras:
        problems.append("unplanned figure legend(s): " + ", ".join(extras))
    return problems


def audit_docx(path: Path, style: dict, kind: str = "",
               planned_figures: list[str] | None = None,
               legend_cap: int = 180) -> list[str]:
    problems: list[str] = []
    if not path.is_file():
        return [f"missing DOCX: {path}"]
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad:
                return [f"{path.name}: corrupt at {bad}"]
            xmls = {n: zf.read(n) for n in zf.namelist()
                    if n.startswith("word/") and n.endswith((".xml", ".rels"))}
    except zipfile.BadZipFile:
        return [f"{path.name}: invalid ZIP/DOCX"]

    document = xmls.get("word/document.xml", b"")
    try:
        root = ET.fromstring(document)
    except ET.ParseError as exc:
        return [f"{path.name}: invalid document.xml ({exc})"]
    if root.findall(".//w:hyperlink", NS):
        problems.append(f"{path.name}: hyperlink elements remain")
    if any("HYPERLINK" in (node.text or "").upper()
           for node in root.findall(".//w:instrText", NS)):
        problems.append(f"{path.name}: hyperlink field code remains")
    for name, blob in xmls.items():
        if name.endswith(".rels"):
            try:
                relroot = ET.fromstring(blob)
            except ET.ParseError:
                continue
            for rel in relroot.findall("rel:Relationship", NS):
                if rel.get("TargetMode") == "External" and "hyperlink" in rel.get("Type", "").lower():
                    problems.append(f"{path.name}: external hyperlink relationship remains")
    for name, blob in xmls.items():
        if not name.endswith(".xml"):
            continue
        part = ET.fromstring(blob)
        if any(node.get(f"{{{W}}}type") in (None, "textWrapping")
               for node in part.findall(".//w:br", NS)) or part.findall(".//w:cr", NS):
            problems.append(f"{path.name}: manual line-break control remains in {name}")
        if any("\u2193" in (node.text or "") for node in part.findall(".//w:t", NS)):
            problems.append(f"{path.name}: forbidden down-arrow character remains in {name}")
        for tag in FORBIDDEN_PARAGRAPH_CONTROLS:
            if part.findall(f".//w:{tag}", NS):
                problems.append(f"{path.name}: forbidden {tag} paragraph control remains in {name}")
        if part.findall(".//w:pBdr", NS):
            problems.append(f"{path.name}: paragraph border/horizontal-rule residue remains in {name}")
    expected_font = str(style["font_family"])
    if expected_font.casefold() == "times new roman" and any(b"Aptos" in b for b in xmls.values()):
        problems.append(f"{path.name}: Aptos remains in Word XML")
    for p in root.findall(".//w:p", NS):
        ppr = p.find("w:pPr", NS)
        pstyle = None if ppr is None else ppr.find("w:pStyle", NS)
        style_id = "" if pstyle is None else pstyle.get(f"{{{W}}}val", "")
        allowed_pts = ({float(style["title_font_pt"])} if style_id == "ManuscriptTitle" else
                       {float(style["section_heading_font_pt"]),
                        float(style["subsection_heading_font_pt"])} if style_id == "SectionHeading" else
                       {float(style["body_font_pt"])})
        paragraph_text = "".join(t.text or "" for t in p.findall(".//w:t", NS)).strip()
        spacing_node = None if ppr is None else ppr.find("w:spacing", NS)
        if paragraph_text and (spacing_node is None or
                spacing_node.get(f"{{{W}}}line") !=
                str(int(round(float(style["line_spacing"]) * 240))) or
                spacing_node.get(f"{{{W}}}after", "0") !=
                str(int(round(float(style["paragraph_spacing_after_pt"]) * 20)))):
            problems.append(f"{path.name}: inconsistent line/paragraph spacing")
        for run in p.findall(".//w:r", NS):
            if not "".join(t.text or "" for t in run.findall(".//w:t", NS)).strip():
                continue
            rpr = run.find("w:rPr", NS)
            rf = None if rpr is None else rpr.find("w:rFonts", NS)
            color = None if rpr is None else rpr.find("w:color", NS)
            size = None if rpr is None else rpr.find("w:sz", NS)
            if rf is None or any(rf.get(f"{{{W}}}{a}") != expected_font for a in ("ascii", "hAnsi")):
                problems.append(f"{path.name}: a text run does not explicitly use {expected_font}")
                break
            if color is None or color.get(f"{{{W}}}val", "").upper() not in {"000000", "AUTO"}:
                problems.append(f"{path.name}: a text run is not explicitly black")
                break
            if size is None or size.get(f"{{{W}}}val", "") not in {
                    str(int(round(value * 2))) for value in allowed_pts}:
                problems.append(f"{path.name}: inconsistent size in {style_id or 'body'} paragraph")
                break
        if style_id.casefold().startswith("heading"):
            problems.append(f"{path.name}: built-in outline heading style remains on visible text")
            break
    text = docx_text(path)
    if "\u2193" in text:
        problems.append(f"{path.name}: forbidden down-arrow character remains")
    if kind == "manuscript":
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        exact = {name: [i for i, line in enumerate(lines) if line.casefold() == name]
                 for name in ("abstract", "references", "figure legends")}
        keywords = [i for i, line in enumerate(lines) if line.casefold().startswith("keywords:")]
        for name, indices in exact.items():
            if len(indices) != 1:
                problems.append(f"{path.name}: expected one {name.title()} heading, found {len(indices)}")
        if len(keywords) != 1:
            problems.append(f"{path.name}: expected one Keywords line, found {len(keywords)}")
        order = [exact[name][0] if len(exact[name]) == 1 else -1
                 for name in ("abstract",)] + [keywords[0] if len(keywords) == 1 else -1] + [
                 exact[name][0] if len(exact[name]) == 1 else -1
                 for name in ("references", "figure legends")]
        if all(x >= 0 for x in order) and order != sorted(order):
            problems.append(f"{path.name}: Abstract/Keywords/References/Figure legends order is wrong")
        if planned_figures is None:
            problems.append(f"{path.name}: artifact_plan.json is missing or invalid")
        else:
            problems.extend(f"{path.name}: {problem}" for problem in
                            _legend_problems(root, planned_figures, legend_cap))
    elif kind == "figure_legends":
        count = sum(1 for line in docx_text(path).splitlines()
                    if line.strip().casefold() == "figure legends")
        if count != 1:
            problems.append(f"{path.name}: expected one Figure legends heading, found {count}")
    if kind == "supplementary" and root.findall(".//w:tbl", NS):
        problems.append(
            f"{path.name}: supplementary Methods contains a table; package it as a separate "
            "three-line supplementary table"
        )
    return sorted(set(problems))


def run_pandoc(inputs: list[Path], output: Path, bibliography: Path | None,
               csl: Path | None) -> None:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("pandoc is required to render Markdown and citations")
    cmd = [pandoc, *map(str, inputs)]
    if bibliography:
        cmd += ["--citeproc", f"--bibliography={bibliography}"]
    if csl:
        cmd += [f"--csl={csl}"]
    cmd += ["--metadata=link-citations:false", "-o", str(output)]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    combined = (proc.stdout or "") + (proc.stderr or "")
    if proc.returncode:
        raise RuntimeError(combined.strip() or f"pandoc exited {proc.returncode}")
    if re.search(r"could not find citation|unresolved", combined, re.I):
        raise RuntimeError("pandoc reported an unresolved citation: " + combined.strip())


def build(args) -> int:
    style = load_style(args.style_config)
    if style.get("all_text_black") is not True or style.get("external_hyperlinks") is not False:
        print("style config must require black text and prohibit external hyperlinks", file=sys.stderr)
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    try:
        source_problems = _source_problems(args.kind, args.input)
        if source_problems:
            raise ValueError("; ".join(source_problems))
        run_pandoc(args.input, args.output, args.bibliography, args.csl)
        normalize_docx(args.output, style, args.kind)
        problems = audit_docx(args.output, style, args.kind,
                              planned_figure_ids() if args.kind == "manuscript" else None,
                              figure_legend_cap())
    except Exception as exc:  # noqa: BLE001
        args.output.unlink(missing_ok=True)
        print(f"DOCX build failed: {exc}", file=sys.stderr)
        return 2
    if problems:
        args.output.unlink(missing_ok=True)
        print("DOCX audit failed:\n  " + "\n  ".join(problems), file=sys.stderr)
        return 2
    print(f"built and audited {args.kind} -> {args.output}")
    return 0


def _resolve_manifest_file(manifest: Path, raw: str) -> Path:
    p = Path(raw)
    if p.is_absolute():
        return p
    project = project_root()
    return project / p


def audit_manifest(args) -> int:
    style = load_style(args.style_config)
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    problems: list[str] = []
    roles: dict[str, list[Path]] = {}
    planned = planned_figure_ids()
    legend_cap = figure_legend_cap()
    for item in manifest.get("items", []):
        role = str(item.get("role", "")).strip()
        path = _resolve_manifest_file(args.manifest, str(item.get("file", "")))
        roles.setdefault(role, []).append(path)
        if role in TEXT_ROLES and path.suffix.lower() != ".docx":
            problems.append(f"{role}: narrative upload must be DOCX, got {path.name}")
        if path.suffix.lower() == ".docx":
            problems += audit_docx(path, style, role, planned, legend_cap)
    for role in ("manuscript", "title_page", "cover_letter"):
        if not any(p.suffix.lower() == ".docx" for p in roles.get(role, [])):
            problems.append(f"manifest lacks a DOCX item for role '{role}'")
    if (project_root() / "08_submission/integration/supplementary_methods.md").exists():
        if not any(p.suffix.lower() == ".docx" for p in roles.get("supplementary", [])):
            problems.append("supplementary_methods.md exists but no supplementary DOCX is listed")
    if problems:
        print("submission DOCX audit failed:\n  " + "\n  ".join(sorted(set(problems))))
        return 2
    print(f"submission DOCX audit passed: {sum(len(v) for v in roles.values())} manifest item(s)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="build or audit deterministic journal DOCX files")
    sub = ap.add_subparsers(dest="command", required=True)
    p = sub.add_parser("build")
    p.add_argument("--kind", required=True,
                   choices=["manuscript", "title_page", "cover_letter", "supplementary",
                            "statements", "checklist_form", "coi_forms", "figure_legends"])
    p.add_argument("--input", type=Path, action="append", required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--style-config", type=Path, required=True)
    p.add_argument("--bibliography", type=Path)
    p.add_argument("--csl", type=Path)
    p.set_defaults(fn=build)
    p = sub.add_parser("audit")
    p.add_argument("--manifest", type=Path, required=True)
    p.add_argument("--style-config", type=Path, required=True)
    p.set_defaults(fn=audit_manifest)
    args = ap.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    raise SystemExit(main())
