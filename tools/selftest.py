#!/usr/bin/env python3
"""End-to-end self test. Builds a throwaway project, exercises the toolchain against the
real gate checks, then deletes everything it made.

    .venv/Scripts/uv run python tools/selftest.py
    .venv/Scripts/uv run python tools/selftest.py --keep     # leave the temp project for inspection
    .venv/Scripts/uv run python tools/selftest.py --online   # also hit PubMed, Crossref

Run this after changing the style module, the table writer, the QC script or any check.
It is the regression test for the parts of the pipeline that code can verify.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

PASS, FAIL = "PASS", "FAIL"
results: list[tuple[str, str, str]] = []


def record(name: str, ok: bool, detail: str = "") -> bool:
    results.append((name, PASS if ok else FAIL, detail))
    print(f"  [{PASS if ok else FAIL}] {name}" + (f"  {detail}" if detail else ""))
    return ok


def section(title: str) -> None:
    print(f"\n{title}\n" + "-" * len(title))


# ---------------------------------------------------------------------------
def build_fixture(proj: Path) -> None:
    """A minimal but structurally complete project: results JSON, tables, a figure,
    an artifact plan, legends, captions and a Results section that cites them."""
    import numpy as np
    from figures.style import PALETTE, apply_style, figure, save, significance
    from tables.threeline import ci, fmt, p_value, write_table, write_workbook

    for sub in ("03_analysis/results", "04_tables/main", "04_tables/supplementary",
                "05_figures/out", "05_figures/qc", "01_protocol", "07_manuscript", "temp"):
        (proj / sub).mkdir(parents=True, exist_ok=True)

    res = {
        "analysis": "primary_model", "script": "03_analysis/code/03_primary.py",
        "n_eligible": 1380, "n_analysed": 1284, "n_excluded": 96,
        "groups": {"a": {"n": 642, "age_mean": 62.14, "age_sd": 11.42},
                   "b": {"n": 642, "age_mean": 61.73, "age_sd": 12.05}},
        "estimates": [{"term": "exposure", "estimate": 1.87, "ci_low": 1.34,
                       "ci_high": 2.61, "p": 0.00021, "scale": "HR"}],
        "age_p": 0.62,
        "roc": {"auc": 0.781, "auc_low": 0.742, "auc_high": 0.820},
    }
    (proj / "03_analysis/results/primary.json").write_text(
        json.dumps(res, indent=2), encoding="utf-8")

    g = res["groups"]
    write_table(
        path=proj / "04_tables/main/Table1.xlsx", sheet="Table 1",
        title="Table 1. Baseline characteristics of the study population.",
        header=["Characteristic", f"Group A (n={g['a']['n']})",
                f"Group B (n={g['b']['n']})", "P value"],
        rows=[["Age, years, mean (SD)",
               f"{fmt(g['a']['age_mean'])} ({fmt(g['a']['age_sd'])})",
               f"{fmt(g['b']['age_mean'])} ({fmt(g['b']['age_sd'])})",
               p_value(res["age_p"])],
              ["Exposure, HR (95% CI)", ci(1.87, 1.34, 2.61), "reference",
               p_value(0.00021)]],
        footnotes=["Data are mean (SD) unless stated otherwise.",
                   "Abbreviations: SD, standard deviation; HR, hazard ratio; "
                   "CI, confidence interval.",
                   "P values from the two-sided t test."],
    )
    write_workbook(
        path=proj / "04_tables/supplementary/supplementary_tables.xlsx",
        tables=[{"sheet": "Table S1", "title": "Table S1. Sensitivity analyses.",
                 "header": ["Model", "HR (95% CI)", "P value"],
                 "rows": [["Complete case", ci(1.87, 1.34, 2.61), p_value(0.00021)]],
                 "footnotes": ["Abbreviations: HR, hazard ratio; CI, confidence interval."]},
                {"sheet": "Table S2", "title": "Table S2. Discrimination.",
                 "header": ["Model", "AUC (95% CI)"],
                 "rows": [["Primary", ci(0.781, 0.742, 0.820, 3)]],
                 "footnotes": ["Abbreviations: AUC, area under the curve."]}],
    )

    apply_style()
    fig, panels = figure(width="double", height_mm=75, panels=(1, 2))
    (sfA, axA), (sfB, axB) = panels
    rng = np.random.default_rng(7)
    fpr = np.linspace(0, 1, 200)
    axA.plot(fpr, fpr ** 0.42, color=PALETTE[0], label=f"Primary (AUC {res['roc']['auc']:.3f})")
    axA.plot([0, 1], [0, 1], ls="--", lw=0.6, color="#4D4D4D")
    axA.set_xlabel("1 - specificity"); axA.set_ylabel("Sensitivity")
    axA.set_xlim(0, 1); axA.set_ylim(0, 1); axA.legend(loc="lower right")
    sfA.suptitle("Discrimination")

    vals = [rng.normal(m, 1.0, 60) for m in (3.1, 4.4)]
    bp = axB.boxplot(vals, widths=0.5, patch_artist=True, showfliers=False)
    for patch, col in zip(bp["boxes"], PALETTE):
        patch.set_facecolor(col); patch.set_alpha(0.35); patch.set_edgecolor("black")
    for key in ("medians", "whiskers", "caps"):
        for art in bp[key]:
            art.set_color("black"); art.set_linewidth(0.6)
    axB.set_xticks([1, 2]); axB.set_xticklabels(["Group A", "Group B"])
    axB.set_ylabel("Grip strength (kg)")
    significance(axB, 1, 2, max(v.max() for v in vals) + 0.3, "***")
    sfB.suptitle("Grip strength by group")
    save(fig, proj / "05_figures/out/Figure1", width="double")

    plan = {
        "main_figures": [{"id": "Figure 1", "slug": "primary",
                          "title": "Discrimination and grip strength",
                          "content": "ROC curve and grip strength by group",
                          "archetype": "other",
                          "archetype_rationale": "composite plate: panel A is a ROC curve, "
                                                 "panel B a box plot; no single archetype fits",
                          "panels": ["A", "B"], "width": "double",
                          "script": "05_figures/code/fig1_primary.py",
                          "file": "05_figures/out/Figure1.png",
                          "tiff": "05_figures/out/Figure1.tiff",
                          "source_results": ["03_analysis/results/primary.json"]}],
        "main_tables": [{"id": "Table 1", "slug": "baseline",
                         "title": "Baseline characteristics", "content": "cohort by group",
                         "file": "04_tables/main/Table1.xlsx", "sheet": "Table 1",
                         "source_results": ["03_analysis/results/primary.json"]}],
        "supp_figures": [],
        "supp_tables": [
            {"id": "Table S1", "slug": "sens", "title": "Sensitivity analyses",
             "content": "robustness",
             "file": "04_tables/supplementary/supplementary_tables.xlsx",
             "sheet": "Table S1", "source_results": ["03_analysis/results/primary.json"]},
            {"id": "Table S2", "slug": "disc", "title": "Discrimination", "content": "AUC",
             "file": "04_tables/supplementary/supplementary_tables.xlsx",
             "sheet": "Table S2", "source_results": ["03_analysis/results/primary.json"]}],
        "supp_files": [],
    }
    (proj / "01_protocol/artifact_plan.json").write_text(
        json.dumps(plan, indent=2), encoding="utf-8")
    (proj / "05_figures/code").mkdir(parents=True, exist_ok=True)
    (proj / "05_figures/code/fig1_primary.py").write_text("# fixture\n", encoding="utf-8")

    (proj / "05_figures/legends.md").write_text(
        "# Figure legends\n\n## Figure 1.\nDiscrimination and grip strength. (A) Receiver "
        "operating characteristic curve for the primary model; the dashed line marks chance "
        "performance. (B) Grip strength by group as box plots (median, interquartile range, "
        "whiskers to 1.5x IQR, outliers omitted). n=1284. ***P<0.001, two-sided Wilcoxon "
        "rank-sum test. Abbreviations: AUC, area under the curve; IQR, interquartile range.\n",
        encoding="utf-8")
    (proj / "04_tables/table_captions.md").write_text(
        "# Table captions\n\n## Table 1.\nBaseline characteristics of the study population. "
        "Data are mean (SD) unless stated otherwise. Abbreviations: SD, standard deviation.\n\n"
        "## Table S1.\nSensitivity analyses. Abbreviations: HR, hazard ratio.\n\n"
        "## Table S2.\nDiscrimination. Abbreviations: AUC, area under the curve.\n",
        encoding="utf-8")
    (proj / "07_manuscript/results.md").write_text(
        "# Results\n\nOf 1380 eligible participants, 96 were excluded and 1284 were "
        "analysed. Mean age was 62.14 (SD 11.42) years in group A and 61.73 (SD 12.05) "
        "years in group B (P=0.62) (Table 1).\n\nExposure was associated with the outcome "
        "(HR 1.87, 95% CI 1.34 to 2.61; P<0.001). The primary model discriminated moderately "
        "(AUC 0.781, 95% CI 0.742 to 0.820) (Figure 1A). Grip strength differed between "
        "groups (Figure 1B). Results were unchanged in the complete-case analysis "
        "(Table S1), and discrimination was similar (Table S2).\n",
        encoding="utf-8")


# ---------------------------------------------------------------------------
def run_checks(proj: Path) -> None:
    from wfcore import gates, registry
    from wfcore.checks import Ctx, Result, get, load_all
    from wfcore.state import State

    load_all()
    pipe = registry.load()
    st = State(proj, ".wf")
    st.create(pipe.meta["name"], pipe.meta["version"], pipe.first().id)

    def run(check_name: str, stage_id: str, **spec) -> Result:
        fn = get(check_name)
        ctx = Ctx(pipeline=pipe, state=st, project=proj,
                  stage=pipe.stage(stage_id), spec={"check": check_name, **spec})
        return fn(ctx)

    section("structural checks (expected to pass)")
    for name, stage, spec in [
        ("tables_threeline", "S10_tables", {}),
        ("tables_match_plan", "S10_tables", {}),
        ("artifact_plan_sane", "S07_artifacts", {"path": "01_protocol/artifact_plan.json"}),
        ("legends_cover_plan", "S07_artifacts", {}),
        ("legend_no_results_restatement", "S07_artifacts", {}),
        ("forbidden_symbols_absent", "S07_artifacts",
         {"paths": ["05_figures/legends.md", "04_tables/table_captions.md"]}),
        ("figures_match_plan", "S11_figures", {}),
        ("numbers_have_provenance", "S10_tables", {"source": "tables"}),
        ("numbers_have_provenance", "S09_results", {"path": "07_manuscript/results.md"}),
        ("artifact_refs_consistent", "S09_results",
         {"paths": ["07_manuscript/results.md"], "require_all_cited": True}),
        ("no_ai_boilerplate", "S09_results", {"path": "07_manuscript/results.md"}),
        ("temp_clean", "S10_tables", {}),
    ]:
        r = run(name, stage, **spec)
        record(f"{name}({spec.get('source') or spec.get('path') or 'plan'})", r.ok, r.detail)

    section("negative controls (checks must catch injected faults)")

    bad = proj / "07_manuscript/bad.md"
    bad.write_text("# Results\n\nThe hazard ratio was 4.44 (95% CI 2.01 to 9.87).\n",
                   encoding="utf-8")
    r = run("numbers_have_provenance", "S09_results", path="07_manuscript/bad.md")
    record("invented statistic is rejected", not r.ok, r.detail[:110])

    bad.write_text("# Results\n\nSee Figure 9 and Table 7 for details.\n", encoding="utf-8")
    r = run("artifact_refs_consistent", "S09_results", paths=["07_manuscript/bad.md"])
    record("citation to an unplanned artifact is rejected", not r.ok, r.detail[:110])

    bad.write_text("# Results\n\nThe adjusted estimate was TODO and warrants attention.\n",
                   encoding="utf-8")
    r = run("no_ai_boilerplate", "S09_results", path="07_manuscript/bad.md")
    record("placeholder text is rejected", not r.ok, r.detail[:110])

    bad.write_text("# Results\n\nWe cite [@fabricated2021smith] here.\n", encoding="utf-8")
    r = run("citekeys_resolve", "S08_methods", paths=["07_manuscript/bad.md"])
    record("citekey with no bib entry is rejected", not r.ok, r.detail[:110])
    bad.unlink()

    legends = proj / "05_figures/legends.md"
    legend_good = legends.read_text(encoding="utf-8")
    legends.write_text("# Figure legends\n\n## Figure 1.\n" +
                       ("Repeated implementation detail that belongs in supplementary Methods. " * 35),
                       encoding="utf-8")
    r = run("legends_cover_plan", "S07_artifacts")
    record("overlong figure legend is rejected", not r.ok, r.detail[:110])
    legends.write_text(legend_good, encoding="utf-8")

    legends.write_text(
        "# Figure legends\n\n## Figure 1.\nExposure was significantly associated with "
        "the outcome (HR 1.87). Panel A displays the model and panel B displays groups.\n",
        encoding="utf-8")
    r = run("legend_no_results_restatement", "S07_artifacts")
    record("figure legend result restatement is rejected", not r.ok, r.detail[:110])
    legends.write_text(
        "# Figure legends\n\n## Figure 1.\nThe lower and upper whiskers represent the "
        "limits of the confidence interval. Error bars show SD or 95% CI. "
        "The reference line marks HR = 1; shading identifies the reference group.\n",
        encoding="utf-8")
    r = run("legend_no_results_restatement", "S07_artifacts")
    record("legend decoding terms are not mistaken for results", r.ok, r.detail[:110])
    legends.write_text(legend_good, encoding="utf-8")

    legends.write_text(legend_good.replace("Discrimination", "\u2193 Discrimination", 1),
                       encoding="utf-8")
    r = run("forbidden_symbols_absent", "S07_artifacts",
            paths=["05_figures/legends.md", "04_tables/table_captions.md"])
    record("literal down arrow is rejected in source", not r.ok, r.detail[:110])
    legends.write_text(legend_good, encoding="utf-8")

    (proj / "03_analysis/code").mkdir(parents=True, exist_ok=True)
    plot = proj / "03_analysis/code/oops.py"
    plot.write_text("import matplotlib.pyplot as plt\nplt.plot([1,2])\n", encoding="utf-8")
    r = run("no_plot_calls", "S05_analysis", glob="03_analysis/code/*.*")
    record("plotting during exploratory analysis is rejected", not r.ok, r.detail[:110])
    plot.unlink()

    scratch = proj / "temp/scratch.csv"
    scratch.write_text("a,b\n1,2\n", encoding="utf-8")
    r = run("temp_clean", "S10_tables")
    record("leftover scratch file is rejected", not r.ok, r.detail[:110])
    scratch.unlink()

    r = run("decision_recorded", "S11_figures",
            name="figures_visually_confirmed", allowed=["YES"])
    record("unrecorded decision is rejected", not r.ok, r.detail[:110])

    st.record_decision("figures_visually_confirmed", "YES", "too short")
    r = run("decision_recorded", "S11_figures",
            name="figures_visually_confirmed", allowed=["YES"])
    record("decision with a thin rationale is rejected", not r.ok, r.detail[:110])

    section("three-line writer rejects malformed input")
    from tables.threeline import write_table
    for label, kwargs in [
        ("missing footnotes", dict(footnotes=[])),
        ("title without a Table N prefix", dict(title="Baseline characteristics")),
        ("prose dumped into a cell", dict(rows=[["x", "y" * 400, "z", "w"]])),
    ]:
        base = dict(path=proj / "temp/bad.xlsx", sheet="Table 9",
                    title="Table 9. Fixture.", header=["a", "b", "c", "d"],
                    rows=[["1", "2", "3", "4"]], footnotes=["Abbreviations: none."])
        base.update(kwargs)
        try:
            write_table(**base)
            record(label, False, "writer accepted it")
        except ValueError as exc:
            record(label, True, str(exc)[:90])
    (proj / "temp/bad.xlsx").unlink(missing_ok=True)

    section("gate runner over every stage (must not raise)")
    raised = []
    for stage in pipe.stages:
        for r in gates.run_stage(pipe, st, proj, stage):
            if "check raised" in r.detail:
                raised.append(f"{stage.id}/{r.check}: {r.detail}")
    record("no check raised an exception", not raised, "; ".join(raised[:3]))

    migration_root = proj / "temp/migration_fixture"
    legacy = State(migration_root, ".wf")
    legacy.create("medpaper", "1.2.0", "S18_journal")
    changed = legacy.migrate_pipeline(
        pipe.meta["version"], [s.id for s in pipe.stages],
        {"S17_frontmatter": "S17_assemble", "S18_journal": "S17_assemble",
         "S19_polish": "S17_assemble", "S20_package": "S17_assemble"})
    record("legacy tail state migrates without resetting project artifacts",
           changed and legacy.current == "S17_assemble" and
           legacy.data["pipeline_version"] == "1.4.0",
           f"current={legacy.current}; version={legacy.data['pipeline_version']}")
    shutil.rmtree(migration_root, ignore_errors=True)

    freeze_migration_root = proj / "temp/freeze_migration_fixture"
    late_legacy = State(freeze_migration_root, ".wf")
    late_legacy.create("medpaper", "1.3.9", "S24_package_human_review")
    migrate_env = {**os.environ, "MEDPAPER_PROJECT": str(freeze_migration_root),
                   "MEDPAPER_ROOT": str(ROOT), "PYTHONIOENCODING": "utf-8"}
    migrated = subprocess.run(
        [sys.executable, str(ROOT / "tools/wf.py"), "status", "--brief"],
        capture_output=True, text=True, env=migrate_env, encoding="utf-8", errors="replace")
    migrated_state = State(freeze_migration_root, ".wf").load()
    record("pre-freeze late-stage workspace migrates back to S19 without deleting artifacts",
           migrated.returncode == 0 and migrated_state.current == "S19_human_review" and
           migrated_state.data["pipeline_version"] == "1.4.0",
           f"current={migrated_state.current}; version={migrated_state.data['pipeline_version']}")
    shutil.rmtree(freeze_migration_root, ignore_errors=True)

    decision_root = proj / "temp/decision_invalidation_fixture"
    revisited = State(decision_root, ".wf")
    revisited.create("medpaper", pipe.meta["version"], "S24_package_human_review")
    revisited.record_decision(
        "submission_package_user_confirmed", "OK",
        "The user reviewed the frozen package and explicitly approved independent audit.")
    revisited.complete("S24_package_human_review", "S25_submission_audit")
    revisited.record_decision(
        "submission_package_independent_audit", "REVISE",
        "The reviewer found a material mismatch that requires a corrected package revision.")
    revisited.clear_decisions_for(["S24_package_human_review", "S25_submission_audit"])
    record("loop invalidates package confirmation and audit decisions",
           revisited.decision("submission_package_user_confirmed") is None and
           revisited.decision("submission_package_independent_audit") is None)
    shutil.rmtree(decision_root, ignore_errors=True)


def run_data_proof(proj: Path) -> None:
    """Full acquisition must be proved; a synchronized partial payload still fails."""
    from wfcore import dataproof, registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    section("data acquisition completeness and anti-truncation")
    scoped = proj / "temp/data_proof"
    raw = scoped / "02_data/raw"
    acquisition = scoped / "02_data/acquisition"
    results_dir = scoped / "03_analysis/results"
    raw.mkdir(parents=True)
    acquisition.mkdir(parents=True)
    results_dir.mkdir(parents=True)
    data_path = raw / "cohort.csv"
    count_path = raw / "source_total.txt"
    code_path = acquisition / "fetch.py"
    data_path.write_text("id,value\n1,10\n2,20\n3,30\n4,40\n5,50\n", encoding="utf-8")
    count_path.write_text("5\n", encoding="utf-8")
    code_path.write_text("# API cursor loop runs until next_cursor is empty\n", encoding="utf-8")

    manifest = {
        "schema_version": dataproof.SCHEMA_VERSION,
        "requested_scope": dataproof.FULL_SCOPE,
        "status": dataproof.COMPLETE,
        "protocol_population": "Every record returned by the pre-specified eligible cohort query.",
        "acquisition_method": "scripted",
        "pilot": {"used": False},
        "protocol_sampling": None,
        "sources": [{
            "id": "eligible_cohort",
            "type": "database",
            "locator": "fixture://eligible-cohort?release=2026-09-01",
            "retrieved_at": "2026-09-07T10:00:00+08:00",
            "source_total_expected": 5,
            "records_requested": 5,
            "records_received": 5,
            "source_total_evidence": {
                "path": "02_data/raw/source_total.txt", "kind": "text_integer"
            },
            "received_count_evidence": [{
                "path": "02_data/raw/cohort.csv", "kind": "delimited_rows", "header": True
            }],
            "complete": True,
            "truncation_applied": False,
            "pagination": {
                "applicable": True, "pages_expected": 2, "pages_received": 2,
                "terminal_reached": True, "next_cursor_at_end": None,
            },
        }],
        "analysis_dataset": {
            "primary_source_id": "eligible_cohort", "n_rows": 4,
            "excluded_after_acquisition": 1,
            "relationship": "one_record_per_analysis_row",
        },
    }

    def synchronize() -> None:
        current_raw = dataproof.inventory(scoped, "02_data/raw")
        manifest["raw_files"] = current_raw
        manifest["acquisition_files"] = dataproof.inventory(scoped, "02_data/acquisition")
        manifest["raw_bundle_sha256"] = dataproof.bundle_sha256(current_raw)
        (scoped / dataproof.MANIFEST_REL).write_text(
            json.dumps(manifest, indent=2), encoding="utf-8")
        summary = {
            "n_rows": manifest["analysis_dataset"]["n_rows"], "n_cols": 2,
            "variables": [], "missingness": {},
            "source_hash": manifest["raw_bundle_sha256"],
            "acquisition_scope": dataproof.FULL_SCOPE,
            "acquisition_status": manifest["status"],
            "built_by": "03_analysis/code/build.py",
        }
        (scoped / dataproof.SUMMARY_REL).write_text(
            json.dumps(summary, indent=2), encoding="utf-8")

    synchronize()
    load_all()
    pipe = registry.load()
    state = State(scoped, ".wf")
    state.create("medpaper", pipe.meta["version"], "S04_data")

    def gate():
        return get("data_acquisition_complete")(Ctx(
            pipeline=pipe, state=state, project=scoped, stage=pipe.stage("S04_data"),
            spec={"check": "data_acquisition_complete"},
        ))

    outcome = gate()
    record("full payload, source total, received count and terminal cursor pass",
           outcome.ok, outcome.detail[:140])

    # Updating hashes and the claimed received count still cannot turn a first-four
    # payload into the independently preserved source total of five.
    data_path.write_text("id,value\n1,10\n2,20\n3,30\n4,40\n", encoding="utf-8")
    manifest["sources"][0]["records_received"] = 4
    manifest["analysis_dataset"]["n_rows"] = 3
    synchronize()
    outcome = gate()
    record("synchronized partial payload is rejected against source-total evidence",
           not outcome.ok and "claims complete" in outcome.detail, outcome.detail[:140])

    data_path.write_text("id,value\n1,10\n2,20\n3,30\n4,40\n5,50\n", encoding="utf-8")
    manifest["sources"][0]["records_received"] = 5
    manifest["analysis_dataset"]["n_rows"] = 4
    code_path.write_text("records = pandas.read_csv(path).head(2)\n", encoding="utf-8")
    synchronize()
    outcome = gate()
    record("unapproved head/sample/LIMIT-style acquisition cap is rejected",
           not outcome.ok and "unapproved acquisition cap" in outcome.detail,
           outcome.detail[:140])

    env = {**os.environ, "MEDPAPER_PROJECT": str(scoped), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    forced = subprocess.run(
        [sys.executable, str(ROOT / "tools/wf.py"), "advance", "--force", "--note",
         "Attempted force advance must remain blocked by incomplete or capped acquisition."],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("--force cannot bypass full-data acquisition proof",
           forced.returncode != 0 and State(scoped, ".wf").load().current == "S04_data",
           ((forced.stdout or "") + (forced.stderr or "")).strip()[:140])

    code_path.write_text(
        "preview = pandas.read_csv(path).head(2)  # MEDPAPER_PILOT_ONLY\n",
        encoding="utf-8")
    manifest["pilot"] = {
        "used": True, "excluded_from_analysis": True,
        "full_acquisition_completed_after_pilot": True,
    }
    synchronize()
    outcome = gate()
    record("non-analytic pilot is allowed only after the full acquisition completes",
           outcome.ok, outcome.detail[:140])

    code_path.write_text(
        "sample = frame.sample(n=5)  # MEDPAPER_PROTOCOL_SAMPLING\n",
        encoding="utf-8")
    manifest["protocol_sampling"] = {
        "pre_specified": True,
        "protocol_path": "01_protocol/protocol_v1.md#population-and-eligibility",
        "explanation": "Probability sampling was fixed before retrieval for the scientific design.",
    }
    synchronize()
    outcome = gate()
    record("agent cannot silently replace a census with protocol sampling",
           not outcome.ok and "protocol_sampling_authorized" in outcome.detail,
           outcome.detail[:140])
    state.record_decision(
        "protocol_sampling_authorized", "YES",
        "The user explicitly approved this pre-specified probability-sampling design and its tradeoffs.")
    outcome = gate()
    record("user-authorized pre-specified scientific sampling can pass",
           outcome.ok, outcome.detail[:140])
    shutil.rmtree(scoped, ignore_errors=True)


def run_reference_proof(proj: Path) -> None:
    """A boolean or forged DOI must not satisfy any citation gate."""
    from unittest.mock import patch

    from wfcore import refproof, registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    section("reference provenance and anti-forgery")
    scoped = proj / "temp/reference_proof"
    refs = scoped / "06_refs"
    cache = refs / "cache"
    manuscript = scoped / "07_manuscript"
    cache.mkdir(parents=True)
    manuscript.mkdir(parents=True)
    key, pmid, real_doi = "smith2025integrity", "40123456", "10.1000/real-doi"
    raw_real = f"""<?xml version="1.0" encoding="UTF-8"?>
<PubmedArticleSet><PubmedArticle><MedlineCitation><PMID>{pmid}</PMID><Article>
<ArticleTitle>Integrity safeguards for clinical reference verification</ArticleTitle>
<Abstract><AbstractText>Reference provenance was evaluated using independent source records.</AbstractText></Abstract>
<AuthorList><Author><LastName>Smith</LastName><ForeName>Alice</ForeName></Author></AuthorList>
<Journal><JournalIssue><PubDate><Year>2025</Year></PubDate></JournalIssue><Title>Journal of Evidence Integrity</Title></Journal>
<ELocationID EIdType="doi">{real_doi}</ELocationID><PublicationTypeList><PublicationType>Journal Article</PublicationType></PublicationTypeList>
</Article></MedlineCitation><PubmedData><ArticleIdList><ArticleId IdType="doi">{real_doi}</ArticleId></ArticleIdList></PubmedData>
</PubmedArticle></PubmedArticleSet>"""
    library = {
        "entries": [{
            "citekey": key, "pmid": pmid, "doi": real_doi,
            "title": "Integrity safeguards for clinical reference verification",
            "abstract": "Reference provenance was evaluated using independent source records.",
            "authors": [{"last": "Smith", "first": "Alice"}],
            "journal": "Journal of Evidence Integrity", "year": "2025",
            "source": "pubmed",
        }]
    }
    library_path = refs / "library.json"
    library_path.write_text(json.dumps(library, indent=2), encoding="utf-8")
    (refs / "refs.bib").write_text(
        f"@article{{{key}, title={{Integrity safeguards for clinical reference verification}}, "
        f"author={{Smith, Alice}}, journal={{Journal of Evidence Integrity}}, year={{2025}}, "
        f"doi={{{real_doi}}}, pmid={{{pmid}}}}}\n", encoding="utf-8")
    cited = manuscript / "proof.md"
    cited.write_text(f"# Introduction\n\nA verified source is cited [@{key}].\n", encoding="utf-8")
    (refs / "verified.json").write_text(
        json.dumps({"records": {key: {"verified": True, "doi": "10.9999/fabricated"}}}, indent=2),
        encoding="utf-8")

    load_all()
    pipe = registry.load()
    state = State(scoped, ".wf")
    state.create("medpaper", pipe.meta["version"], "S13_reflib")

    def gate(name: str, **spec):
        return get(name)(Ctx(pipeline=pipe, state=state, project=scoped,
                             stage=pipe.stage("S13_reflib"), spec={"check": name, **spec}))

    outcome = gate("citekeys_resolve", paths=["07_manuscript/proof.md"])
    record("verified=true without raw PubMed proof is rejected", not outcome.ok,
           outcome.detail[:120])
    force_env = {**os.environ, "MEDPAPER_PROJECT": str(scoped),
                 "MEDPAPER_ROOT": str(ROOT), "PYTHONIOENCODING": "utf-8"}
    forced = subprocess.run(
        [sys.executable, str(ROOT / "tools/wf.py"), "advance", "--force", "--note",
         "Attempted forced advance must remain blocked by bibliographic provenance."],
        capture_output=True, text=True, env=force_env, encoding="utf-8", errors="replace")
    record("--force cannot bypass reference-integrity gates",
           forced.returncode != 0 and State(scoped, ".wf").load().current == "S13_reflib",
           ((forced.stdout or "") + (forced.stderr or "")).strip()[:120])

    cache_rel = "06_refs/cache/verify_efetch_pubmed_fixture.xml"
    cache_path = scoped / cache_rel
    cache_path.write_text(raw_real, encoding="utf-8")
    source = refproof.parse_pubmed_payload(raw_real)[pmid]
    verified = {
        "schema": refproof.SCHEMA,
        "generator": {"tool": refproof.GENERATOR, "mode": "fresh_ncbi_pubmed_efetch",
                      "source": "NCBI PubMed EFetch"},
        "verified_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "strict": False,
        "library_sha256": refproof.file_sha256(library_path),
        "n_entries": 1, "n_verified": 1,
        "records": {key: {
            "verified": True, "via": "pubmed", "pmid": pmid, "doi": real_doi,
            "evidence": {"source": refproof.SOURCE, "cache_file": cache_rel,
                         "payload_sha256": refproof.file_sha256(cache_path),
                         "record_sha256": refproof.record_sha256(source),
                         "fresh_fetch": True},
        }},
    }
    verified_path = refs / "verified.json"
    verified_path.write_text(json.dumps(verified, indent=2), encoding="utf-8")
    outcome = gate("citekeys_resolve", paths=["07_manuscript/proof.md"])
    record("hashed reparsed PubMed proof satisfies citekey gate", outcome.ok, outcome.detail[:120])

    library["entries"][0]["doi"] = "10.9999/fabricated"
    library_path.write_text(json.dumps(library, indent=2), encoding="utf-8")
    verified["library_sha256"] = refproof.file_sha256(library_path)
    verified["records"][key]["doi"] = "10.9999/fabricated"
    verified_path.write_text(json.dumps(verified, indent=2), encoding="utf-8")
    outcome = gate("citekeys_resolve", paths=["07_manuscript/proof.md"])
    record("forged DOI fails even after booleans and library hash are patched", not outcome.ok,
           outcome.detail[:120])

    # Even a fully forged local XML/hash bundle cannot pass S13: that gate performs
    # its own fresh PubMed fetch and compares the authoritative response.
    raw_fake = raw_real.replace(real_doi, "10.9999/fabricated")
    cache_path.write_text(raw_fake, encoding="utf-8")
    forged_source = refproof.parse_pubmed_payload(raw_fake)[pmid]
    evidence = verified["records"][key]["evidence"]
    evidence["payload_sha256"] = refproof.file_sha256(cache_path)
    evidence["record_sha256"] = refproof.record_sha256(forged_source)
    verified_path.write_text(json.dumps(verified, indent=2), encoding="utf-8")
    live_record = {
        "pmid": pmid, "doi": real_doi,
        "title": "Integrity safeguards for clinical reference verification",
        "abstract": "Reference provenance was evaluated using independent source records.",
        "authors": [{"last": "Smith", "first": "Alice"}],
        "journal": "Journal of Evidence Integrity", "year": "2025", "flags": [],
    }
    with patch("pubmed.eutils.efetch", return_value=[live_record]):
        outcome = gate("reference_provenance", live=True)
    record("independent live S13 gate rejects a locally forged XML receipt", not outcome.ok,
           outcome.detail[:120])

    shutil.rmtree(scoped, ignore_errors=True)


def run_qc(proj: Path) -> None:
    section("figure QC (subprocess, as the agent runs it)")
    env = {**os.environ, "MEDPAPER_PROJECT": str(proj), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    p = subprocess.run([sys.executable, str(ROOT / "tools/figures/qc.py"), "--all"],
                       capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    out = (p.stdout or "") + (p.stderr or "")
    for line in out.splitlines():
        if line.strip().startswith("[") or "pass deterministic" in line:
            print("  " + line.strip())
    record("all deterministic figure QC checks pass", "1/1 figure(s) pass" in out,
           "" if "1/1 figure(s) pass" in out else "see output above")
    rep = proj / "05_figures/qc/qc_report.json"
    record("qc_report.json written", rep.exists())
    if rep.exists():
        data = json.loads(rep.read_text(encoding="utf-8"))
        fig = data["figures"][0]
        record("visual_reviewed defaults to false", fig.get("visual_reviewed") is False,
               "QC alone must never satisfy the visual-verification gate")


SLOP = """# Discussion

In today's rapidly evolving landscape of modern clinical care, sarcopenia plays a crucial
role in patient outcomes. It is worth noting that we utilized a robust and comprehensive
approach to delve into this multifaceted problem, and our cutting-edge analysis paves the
way for a comprehensive understanding of the intricate interplay involved.

Our primary finding was an association between exposure and the outcome (HR 1.87, 95% CI
1.34 to 2.61; P<0.001), which aligns with previous work [@smith2020cohort]. Furthermore,
the effect was seamlessly consistent across strata. Moreover, discrimination was moderate
(AUC 0.781) (Figure 1A). Additionally, results held in sensitivity analyses (Table S1).

The data was analyzed in 3 centers, and 10-20 participants per site were characterised
using a 5mg dose. P-value thresholds were applied and P = 0.000 was observed in one
subgroup. This finding may potentially suggest that the exposure could possibly cause the
outcome, and underscores the importance of further work. Further studies are warranted.
"""

CLEAN = """# Discussion

Sarcopenia was associated with mortality in this cohort. We analysed a single prospective
cohort with prespecified exposure and outcome definitions.

The primary finding was an association between exposure and the outcome (HR 1.87, 95% CI
1.34 to 2.61; P<0.001), consistent with the earlier cohort of Smith and colleagues
[@smith2020cohort]. The effect was stable across strata, discrimination was moderate
(AUC 0.781) (Figure 1A), and the estimate was unchanged in sensitivity analyses (Table S1).

Data were collected at 3 centres, with 10\u201320 participants per site characterised after a
5 mg dose. One subgroup reached P < 0.001. Because the design is observational, these data
support an association rather than a causal effect; the direction of residual confounding
cannot be established from these data.
"""


def run_archetypes(proj: Path) -> None:
    """The archetype gate must detect mandatory elements, not take the plan's word for it."""
    import tomllib

    import matplotlib.pyplot as plt
    import numpy as np
    from figures import elements as el
    from figures.style import apply_style, figure, save
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    section("archetypes: registry")
    reg = tomllib.loads((ROOT / "reference/archetypes.toml").read_text(encoding="utf-8"))
    arches = reg.get("archetype", {})
    record("registry parses", bool(arches), f"{len(arches)} archetype(s)")
    unknown = sorted({k for a in arches.values()
                      for k in list(a.get("requires", [])) + list(a.get("forbids", []))}
                     - set(el.DETECTORS))
    record("every requires/forbids element has a detector", not unknown,
           f"undetectable: {unknown}" if unknown else
           f"{len(el.DETECTORS)} detectors cover all mandatory elements")
    universal = reg.get("meta", {}).get("universal_requires", [])
    record("universal requires are detectable",
           all(u in el.DETECTORS for u in universal), f"{universal}")

    section("archetypes: element detection on constructed figures")
    apply_style()
    rng = np.random.default_rng(3)

    # a compliant bar_dot: points, error bars, baseline at zero
    fig, panels = figure(width="single", height_mm=60, panels=(1, 1), letters=False)
    _, ax = panels[0]
    ax.bar([1, 2], [3.0, 4.2], yerr=[0.4, 0.5], capsize=2.5, color=["C0", "C1"], alpha=0.35,
           edgecolor="black", linewidth=0.6)
    for i, m in ((1, 3.0), (2, 4.2)):
        ax.scatter(np.full(5, i) + rng.normal(0, 0.05, 5), rng.normal(m, 0.3, 5),
                   s=6, color="black", zorder=3)
    ax.set_ylim(0, 5.5)
    ax.set_xticks([1, 2]); ax.set_xticklabels(["Control", "Treated"])
    ax.set_xlabel("Group"); ax.set_ylabel("Relative expression")
    found = el.detect_all(fig)
    for key in reg["archetype"]["bar_dot"]["requires"]:
        record(f"bar_dot detects {key}", found[key]["found"], found[key]["evidence"][:70])
    plt.close(fig)

    # the same figure with a truncated baseline must fail
    fig, panels = figure(width="single", height_mm=60, panels=(1, 1), letters=False)
    _, ax = panels[0]
    ax.bar([1, 2], [3.0, 4.2], color="C0")
    ax.set_ylim(2.5, 4.5)
    ax.set_xlabel("Group"); ax.set_ylabel("Relative expression")
    found = el.detect_all(fig)
    record("truncated bar baseline is detected", not found["baseline_zero"]["found"],
           found["baseline_zero"]["evidence"][:70])
    record("missing individual points are detected", not found["individual_points"]["found"])
    plt.close(fig)

    # missing axis labels
    fig, panels = figure(width="single", height_mm=50, panels=(1, 1), letters=False)
    _, ax = panels[0]
    ax.plot([0, 1], [0, 1])
    found = el.detect_all(fig)
    record("missing axis labels are detected", not found["axis_labels"]["found"],
           found["axis_labels"]["evidence"][:70])
    plt.close(fig)

    # a missing glyph must be caught during rasterization
    fig, panels = figure(width="single", height_mm=40, panels=(1, 1), letters=False)
    _, ax = panels[0]
    ax.plot([0, 1], [0, 1]); ax.set_xlabel("x"); ax.set_ylabel("y")
    ax.text(0.4, 0.5, "\u2f00\u2f01\u4e2d\u6587")     # glyphs absent from the Latin stack
    warns = el.glyph_warnings(fig)
    record("missing font glyphs are caught", bool(warns),
           (warns[0][:70] if warns else "no warning raised - font may cover these glyphs"))
    plt.close(fig)

    # tick labels forced to collide
    fig, panels = figure(width="single", height_mm=40, panels=(1, 1), letters=False)
    _, ax = panels[0]
    ax.plot(range(40), range(40))
    ax.set_xticks(range(40))
    ax.set_xticklabels([f"label{i}" for i in range(40)])
    ax.set_xlabel("x"); ax.set_ylabel("y")
    fig.canvas.draw()
    coll = el.tick_label_collisions(fig, fig.canvas.get_renderer())
    record("tick-label collisions are caught", bool(coll),
           f"{len(coll)} collision(s)" if coll else "none detected")
    plt.close(fig)

    # a deliberate dead band between panels
    fig, panels = figure(width="single", height_mm=90, panels=(1, 1), letters=False)
    sf, ph = panels[0]
    ph.remove()
    a1, a2 = sf.subplots(2, 1, gridspec_kw={"hspace": 1.4})
    for a in (a1, a2):
        a.plot([0, 1], [0, 1]); a.set_xlabel("x"); a.set_ylabel("y")
    fig.canvas.draw()
    voids = el.interior_voids(fig, fig.canvas.get_renderer())
    record("interior dead band is caught", bool(voids),
           f"{voids[0]['gap_pct']}% gap" if voids else "none detected")
    plt.close(fig)

    section("archetypes: gate rejects a plan without an archetype")
    load_all()
    pipe = registry.load()
    st = State(proj, ".wf").load()
    plan_path = proj / "01_protocol/artifact_plan.json"
    original = plan_path.read_text(encoding="utf-8")
    plan = json.loads(original)

    def plan_check():
        fn = get("artifact_plan_sane")
        return fn(Ctx(pipeline=pipe, state=st, project=proj, stage=pipe.stage("S07_artifacts"),
                      spec={"check": "artifact_plan_sane", "path": "01_protocol/artifact_plan.json"}))

    plan["main_figures"][0].pop("archetype", None)
    plan_path.write_text(json.dumps(plan, indent=2), encoding="utf-8")
    r = plan_check()
    record("figure without an archetype is rejected", not r.ok, r.detail[:90])

    plan["main_figures"][0]["archetype"] = "not_a_real_archetype"
    plan_path.write_text(json.dumps(plan, indent=2), encoding="utf-8")
    r = plan_check()
    record("unknown archetype is rejected", not r.ok, r.detail[:90])

    plan["main_figures"][0]["archetype"] = "other"
    plan["main_figures"][0].pop("archetype_rationale", None)
    plan_path.write_text(json.dumps(plan, indent=2), encoding="utf-8")
    r = plan_check()
    record("archetype 'other' without a rationale is rejected", not r.ok, r.detail[:90])

    plan["main_figures"][0]["archetype"] = "roc_curve"
    plan_path.write_text(json.dumps(plan, indent=2), encoding="utf-8")
    r = plan_check()
    record("valid archetype is accepted", r.ok, r.detail[:90])
    plan_path.write_text(original, encoding="utf-8")

    section("archetypes: journal palettes")
    for name in ("okabe_ito", "nejm", "lancet", "jama", "jco", "nature"):
        colors = apply_style(palette=name)
        ok = isinstance(colors, list) and len(colors) >= 5 and all(
            c.startswith("#") and len(c) == 7 for c in colors)
        record(f"palette {name}", ok, f"{len(colors)} colours")
    try:
        apply_style(palette="nope")
        record("unknown palette is rejected", False, "accepted silently")
    except ValueError as exc:
        record("unknown palette is rejected", True, str(exc)[:60])
    apply_style()
    import matplotlib
    record("axes.unicode_minus disabled", matplotlib.rcParams["axes.unicode_minus"] is False,
           "a negative sign cannot render as a missing-glyph box")


def run_codex_integration() -> None:
    """Repository-local Codex integration must be complete and internally consistent."""
    import tomllib

    section("Codex integration")
    required = [
        ROOT / "AGENTS.md",
        ROOT / ".agents/skills/medpaper-pipeline/SKILL.md",
        ROOT / ".agents/skills/medpaper-codex-pipeline/agents/openai.yaml",
        ROOT / "reference/codex-integration.md",
        ROOT / "reference/methods-structure.md",
        ROOT / "reference/rework-routing.md",
        ROOT / "tools/rework.py",
        ROOT / "tools/manuscript/review_package.py",
        ROOT / "tools/wfcore/reviewpackage.py",
        ROOT / "tools/manuscript/scientific_freeze.py",
        ROOT / "tools/wfcore/scientificfreeze.py",
        ROOT / "tools/manuscript/journal_workspace.py",
        ROOT / "tools/wfcore/journalworkspace.py",
        ROOT / "tools/package_content.py",
        ROOT / "tools/wfcore/refproof.py",
        ROOT / "tools/wfcore/dataproof.py",
        ROOT / "tools/data_manifest.py",
        ROOT / "reference/data-acquisition-integrity.md",
    ]
    record("repository Codex files exist", all(p.is_file() for p in required),
           ", ".join(str(p.relative_to(ROOT)) for p in required if not p.is_file()))

    skill = required[1].read_text(encoding="utf-8")
    metadata = required[2].read_text(encoding="utf-8")
    pipe = tomllib.loads((ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8"))
    record("skill identity is valid",
           (skill.startswith("---\nname: medpaper-codex-pipeline\n") or skill.startswith("---\nname: medpaper-pipeline\n")) and
           "description:" in skill.split("---", 2)[1],
           "frontmatter name/description")
    record("skill launcher names the skill",
           'default_prompt: "Use $medpaper-codex-pipeline' in metadata and
           "allow_implicit_invocation: false" in metadata,
           "agents/openai.yaml")
    record("pipeline and skill versions agree",
           pipe["meta"]["version"] == "1.4.0" and 'version: "1.4.0"' in skill,
           f'pipeline={pipe["meta"]["version"]}')
    record("all 25 stage cards exist",
           len(pipe["stage"]) == 25 and
           all((ROOT / "pipeline" / s["card"]).is_file() for s in pipe["stage"]),
           f'{len(pipe["stage"])} stage(s)')

    stage_ids = [s["id"] for s in pipe["stage"]]
    record("manuscript review precedes authors and packaging",
           stage_ids.index("S17_assemble") < stage_ids.index("S18_independent_review") <
           stage_ids.index("S19_human_review") < stage_ids.index("S20_journal") <
           stage_ids.index("S21_authors") < stage_ids.index("S23_package") <
           stage_ids.index("S24_package_human_review") < stage_ids.index("S25_submission_audit"),
           "assemble -> publishability review -> user manuscript review -> journal -> authors -> package -> user package OK -> final audit")
    independent = next(s for s in pipe["stage"] if s["id"] == "S18_independent_review")
    human = next(s for s in pipe["stage"] if s["id"] == "S19_human_review")
    authors = next(s for s in pipe["stage"] if s["id"] == "S21_authors")
    package_human = next(s for s in pipe["stage"] if s["id"] == "S24_package_human_review")
    final_audit = next(s for s in pipe["stage"] if s["id"] == "S25_submission_audit")
    review_card = (ROOT / "pipeline/stages/S18_independent_review.md").read_text(encoding="utf-8")
    package_card = (ROOT / "pipeline/stages/S23_package.md").read_text(encoding="utf-8")
    package_human_card = (ROOT / "pipeline/stages/S24_package_human_review.md").read_text(encoding="utf-8")
    human_card = (ROOT / "pipeline/stages/S19_human_review.md").read_text(encoding="utf-8")
    final_audit_card = (ROOT / "pipeline/stages/S25_submission_audit.md").read_text(encoding="utf-8")
    record("independent reviewer is singular, read-only, and before author intake",
           independent.get("needs_user", False) is False and human.get("needs_user") is True and
           authors.get("needs_user") is True and "exactly one independent subagent" in review_card and
           "must not edit files" in review_card and "Do not ask for authors" in
           (ROOT / "pipeline/stages/S19_human_review.md").read_text(encoding="utf-8"),
           "one frozen-package review, then user review, then authors")
    record("final package audit waits for explicit user OK and is independent",
           package_human.get("needs_user") is True and
           final_audit.get("needs_user", False) is False and
           "是否确认 OK" in package_human_card and
           "tools/package_review.py freeze" in package_human_card and
           "exactly one independent read-only subagent" in final_audit_card and
           "Freeze ID: <freeze_id>" in final_audit_card and
           "allowed = [\"PASS\"]" in
           (ROOT / "pipeline/pipeline.toml").read_text(encoding="utf-8") and
           "loop S24_package_human_review" in final_audit_card,
           "user confirmation -> immutable freeze -> one guideline audit -> PASS or S24 loop")
    final_headings = next(g for g in final_audit["gate"]
                          if g["check"] == "md_sections")["headings"]
    record("final independent audit includes reader and editor lenses",
           "ordinary scientific reader" in final_audit_card and
           "Editorial screening" in final_audit_card and
           "low-level errors" in final_audit_card and
           "Reader comprehension" in final_headings and
           "Editorial screening and low-level errors" in final_headings and
           "objective low-level error" in final_audit_card and
           "never edits the frozen package" in final_audit_card,
           "first-time comprehension + editorial error screen + journal compliance")
    human_checks = {gate["check"] for gate in human["gate"]}
    package_checks = {gate["check"] for gate in package_human["gate"]}
    package_stage = next(s for s in pipe["stage"] if s["id"] == "S23_package")
    record("user revisions are source-routed and mechanically gated",
           "assembly_matches_sources" in human_checks and
           "package_content_matches_baseline" in package_checks and
           "revision_rounds_closed" in human_checks and
           "revision_rounds_closed" in package_checks and
           any(gate["check"] == "revision_rounds_closed" for gate in final_audit["gate"]) and
           "08_submission/package_content_baseline.json" in package_stage["outputs"] and
           "tools/rework.py batch" in package_human_card and
           "tools/rework.py status" in package_human_card and
           "Never recapture at S24" in package_human_card,
           "persisted S19/S24 rounds + component equality + visible-text enforcement")
    record("S19 requires a current third-party ZIP and explicit no-further-review decision",
           "manuscript_review_package_current" in human_checks and
           "s19_review_release_explicit" in human_checks and
           "tools/manuscript/review_package.py build" in human_card and
           "无需继续审核，可以进入下一步" in human_card and
           "05_figures/out/" in human_card and "exact ZIP" in human_card and
           '"s19_review_release_explicit"' in
           (ROOT / "tools/wfcore/cli.py").read_text(encoding="utf-8"),
           "versioned review ZIP + exact paths + package-bound non-overridable stop")
    journal = next(s for s in pipe["stage"] if s["id"] == "S20_journal")
    late_stages = [next(s for s in pipe["stage"] if s["id"] == stage_id) for stage_id in
                   ("S20_journal", "S21_authors", "S22_polish", "S23_package",
                    "S24_package_human_review", "S25_submission_audit")]
    cli_text = (ROOT / "tools/wfcore/cli.py").read_text(encoding="utf-8")
    record("S19 scientific master is frozen and journal adaptation is isolated",
           "07_manuscript/scientific_master_freeze.json" in human["outputs"] and
           any(g["check"] == "scientific_master_frozen" for g in human["gate"]) and
           "08_submission/integration/journal_workspace.json" in journal["outputs"] and
           all(any(g["check"] == "scientific_master_unchanged" for g in stage["gate"])
               for stage in late_stages) and
           all(any(g["check"] == "journal_workspace_ready" for g in stage["gate"])
               for stage in late_stages) and
           all(name in cli_text for name in
               ("scientific_master_frozen", "scientific_master_unchanged",
                "journal_workspace_ready")) and
           "tools/manuscript/scientific_freeze.py freeze" in human_card and
           "tools/manuscript/journal_workspace.py init" in
           (ROOT / "pipeline/stages/S20_journal.md").read_text(encoding="utf-8"),
           "S19 hash freeze -> immutable master -> per-journal integration sources")
    record("Word package contract covers reported defects",
           all(term in package_card for term in ("all text black", "external hyperlinks",
                                                  "supplementary Methods", "Figure legends",
                                                  "collapsible heading")),
           "typography, links, supplementary Word file, legends and heading behavior")
    reflib = next(s for s in pipe["stage"] if s["id"] == "S13_reflib")
    live_reference_stages = {
        stage["id"] for stage in pipe["stage"]
        if any(gate.get("check") == "reference_provenance" and gate.get("live") is True
               for gate in stage.get("gate", []))
    }
    required_live_stages = {
        "S13_reflib", "S14_introduction", "S16_discussion", "S17_assemble",
        "S18_independent_review", "S19_human_review", "S21_authors", "S22_polish",
        "S23_package", "S24_package_human_review", "S25_submission_audit",
    }
    record("reference verification is live, evidence-bound and non-overridable",
           any(g.get("check") == "reference_provenance" and g.get("live") is True
               for g in reflib["gate"]) and
           required_live_stages.issubset(live_reference_stages) and
           "verified: true` by itself has no authority" in
           (ROOT / "pipeline/stages/S13_reflib.md").read_text(encoding="utf-8") and
           "NON_OVERRIDABLE_GATES" in
           (ROOT / "tools/wfcore/cli.py").read_text(encoding="utf-8"),
           "fresh EFetch + raw XML/library hashes + repeated live checks through final audit")
    acquisition_stages = {
        stage["id"] for stage in pipe["stage"]
        if any(gate.get("check") == "data_acquisition_complete"
               for gate in stage.get("gate", []))
    }
    record("full data acquisition is evidence-bound and non-overridable",
           {"S04_data", "S05_analysis", "S06_protocol_final", "S17_assemble",
            "S23_package", "S25_submission_audit"}.issubset(acquisition_stages) and
           "02_data/acquisition_manifest.json" in
           next(s for s in pipe["stage"] if s["id"] == "S04_data")["outputs"] and
           '"data_acquisition_complete"' in
           (ROOT / "tools/wfcore/cli.py").read_text(encoding="utf-8") and
           "full_protocol_defined_universe" in
           (ROOT / "pipeline/stages/S04_data.md").read_text(encoding="utf-8"),
           "source total + received counts + pagination + raw/code hashes + anti-cap scan")

    legacy = [
        ROOT / ".agents/AGENTS.md",
        ROOT / ".medpaper-target",
        ROOT / "tools/install_adapters.py",
        ROOT / "tools/hooks/skill_guard.py",
        ROOT / "reference/skill_policy.toml",
    ]
    record("obsolete multi-agent adapters are absent", not any(p.exists() for p in legacy),
           ", ".join(str(p.relative_to(ROOT)) for p in legacy if p.exists()))

    policy_text = "\n".join(p.read_text(encoding="utf-8") for p in required)
    cards_text = "\n".join(
        (ROOT / "pipeline" / s["card"]).read_text(encoding="utf-8") for s in pipe["stage"])
    record("AI disclosure remains user-controlled",
           "Do not draft or insert an AI-use disclosure" in
           (ROOT / "AGENTS.md").read_text(encoding="utf-8") and
           "Disclose AI assistance honestly" not in cards_text,
           "no autonomous disclosure instruction")
    record("capability arbitration is documented",
           "sole workflow orchestrator" in policy_text and
           "must be registered" in policy_text.lower(),
           "pipeline owns state/gates; helpers remain stage-scoped")

    methods_stage = next(s for s in pipe["stage"] if s["id"] == "S08_methods")
    outputs_gate = next(g for g in methods_stage["gate"] if g["check"] == "outputs_exist")
    methods_card = (ROOT / "pipeline/stages/S08_methods.md").read_text(encoding="utf-8")
    methods_guide = (ROOT / "reference/methods-structure.md").read_text(encoding="utf-8")
    record("supplementary Methods is declared but optional",
           "07_manuscript/supplementary_methods.md" in methods_stage["outputs"] and
           outputs_gate.get("optional") == ["07_manuscript/supplementary_methods.md"] and
           any(g.get("path") == "07_manuscript/supplementary_methods.md" and
               g.get("optional") is True for g in methods_stage["gate"]) and
           not any("exactly ONE manuscript section file" in rule
                   for rule in pipe.get("policy", {}).get("invariants", [])),
           "S08 optional output contract")
    record("Methods architecture prevents subsection inflation",
           "Do not append automatic standalone headings" in methods_card and
           "Statistical analysis" in methods_guide and
           "not a mandatory template" in methods_guide,
           "flexible core structure plus supplement boundary")
    record("supplementary Methods cannot absorb display tables or Markdown rules",
           any(g.get("check") == "supplementary_methods_clean"
               for g in methods_stage["gate"]) and
           "three-line supplementary workbook" in methods_guide and
           "do not embed Markdown" in methods_card,
           "supplement prose only; Table S* stays in the S10 workbook")
    record("ImageGen is review-only for scientific figures",
           "second visual critic" in policy_text and
           "must not redraw" in policy_text,
           "S11 preserves data geometry and requires a fresh render")

    from rework import resolve_route
    from wfcore import registry
    route_pipe = registry.load()
    record("revision router chooses the earliest owning stage",
           resolve_route("analysis", "S24_package_human_review", route_pipe) == "S05_analysis" and
           resolve_route("references", "S19_human_review", route_pipe) == "S13_reflib" and
           resolve_route("discussion", "S19_human_review", route_pipe) == "S16_discussion" and
           resolve_route("manuscript-copyedit", "S19_human_review", route_pipe) ==
           "S19_human_review" and
           resolve_route("manuscript-copyedit", "S24_package_human_review", route_pipe) ==
           "S22_polish" and
           resolve_route("word-format-only", "S25_submission_audit", route_pipe) ==
           "S24_package_human_review",
           "analysis/reference/discussion/copyedit/Word-format routes")

    from wfcore.state import State
    owned_root = os.environ.get("MEDPAPER_SELFTEST_ROOT")
    if owned_root:
        route_root = Path(owned_root) / "rework_route"
        route_root.mkdir(parents=True, exist_ok=False)
    else:
        route_root = Path(tempfile.mkdtemp(prefix="medpaper_rework_route_"))
    route_project = route_root / "project"
    route_state = State(route_project, ".wf")
    route_state.create("medpaper", route_pipe.meta["version"], "S24_package_human_review")
    route_state.record_decision(
        "submission_package_user_confirmed", "OK",
        "The fixture user approved the package before requesting a substantive revision.")
    route_env = {**os.environ, "MEDPAPER_PROJECT": str(route_project),
                 "MEDPAPER_ROOT": str(ROOT), "PYTHONIOENCODING": "utf-8"}
    routed = subprocess.run(
        [sys.executable, str(ROOT / "tools/rework.py"), "start", "--kind", "discussion",
         "--why", "Revise the interpretation requested by the user in the Discussion section.",
         "--project", str(route_project)], capture_output=True, text=True, env=route_env,
        encoding="utf-8", errors="replace")
    routed_state = State(route_project, ".wf").load()
    record("revision router rewinds state and invalidates downstream approval",
           routed.returncode == 0 and routed_state.current == "S16_discussion" and
           routed_state.decision("submission_package_user_confirmed") is None,
           ((routed.stdout or "") + (routed.stderr or "")).strip()[:140])
    shutil.rmtree(route_root, ignore_errors=True)

    if owned_root:
        batch_root = Path(owned_root) / "revision_batch"
        batch_root.mkdir(parents=True, exist_ok=False)
    else:
        batch_root = Path(tempfile.mkdtemp(prefix="medpaper_revision_batch_"))
    batch_project = batch_root / "project"
    manuscript = batch_project / "07_manuscript"
    manuscript.mkdir(parents=True)
    (manuscript / "methods.md").write_text("# Methods\n\nOriginal methods source.\n", encoding="utf-8")
    (manuscript / "discussion.md").write_text("# Discussion\n\nOriginal discussion source.\n", encoding="utf-8")
    batch_state = State(batch_project, ".wf")
    batch_state.create("medpaper", route_pipe.meta["version"], "S24_package_human_review")
    batch_state.record_decision(
        "manuscript_human_reviewed", "NO_FURTHER_REVIEW",
        "The scientific manuscript was approved before journal-specific package revision.")
    batch_state.record_decision(
        "submission_package_user_confirmed", "OK",
        "The package was approved before the user submitted a new multi-item feedback batch.")
    plan_path = batch_root / "revision_plan.json"
    plan_path.write_text(json.dumps({
        "feedback_verbatim": "Please shorten Methods and clarify the causal interpretation in Discussion.",
        "interpretation": "Revise both canonical manuscript sources without changing unsupported facts or bypassing downstream validation.",
        "ambiguous_or_requires_user_decision": [],
        "items": [
            {
                "kind": "methods",
                "request": "Shorten the Methods while preserving every necessary reproducibility detail.",
                "affected_sources": ["07_manuscript/methods.md"],
                "acceptance_criteria": ["Main Methods is concise and complete", "No supported detail is lost"],
            },
            {
                "kind": "discussion",
                "request": "Clarify the causal interpretation and keep claims within the study design.",
                "affected_sources": ["07_manuscript/discussion.md"],
                "acceptance_criteria": ["Claims match the design limitations", "Interpretation is clear to a reader"],
            },
        ],
    }, indent=2), encoding="utf-8")
    batch_env = {**os.environ, "MEDPAPER_PROJECT": str(batch_project),
                 "MEDPAPER_ROOT": str(ROOT), "PYTHONIOENCODING": "utf-8"}
    opened = subprocess.run(
        [sys.executable, str(ROOT / "tools/rework.py"), "batch", "--plan", str(plan_path),
         "--project", str(batch_project)], capture_output=True, text=True, env=batch_env,
        encoding="utf-8", errors="replace")
    batch_state = State(batch_project, ".wf").load()
    round_path = batch_project / ".wf/revisions/R001.json"
    round_payload = json.loads(round_path.read_text(encoding="utf-8")) if round_path.exists() else {}
    record("multi-item revision persists interpretation and rewinds once to earliest owner",
           opened.returncode == 0 and batch_state.current == "S08_methods" and
           batch_state.data.get("active_revision_round") == "R001" and
           round_payload.get("earliest_stage") == "S08_methods" and
           len(round_payload.get("items", [])) == 2 and
           batch_state.decision("manuscript_human_reviewed") is None and
           batch_state.decision("submission_package_user_confirmed") is None,
           ((opened.stdout or "") + (opened.stderr or "")).strip()[:180])
    status = subprocess.run(
        [sys.executable, str(ROOT / "tools/rework.py"), "status", "--project", str(batch_project)],
        capture_output=True, text=True, env=batch_env, encoding="utf-8", errors="replace")
    record("revision status restores atomic work after context compaction",
           status.returncode == 0 and "R001-01" in status.stdout and "R001-02" in status.stdout and
           "acceptance:" in status.stdout and "sources:" in status.stdout,
           status.stdout.strip()[:160])

    from wfcore.checks import Ctx, get, load_all
    load_all()
    revision_gate = get("revision_rounds_closed")
    review_stage = route_pipe.stage("S24_package_human_review")
    blocked = revision_gate(Ctx(route_pipe, batch_state, batch_project, review_stage, {}))
    record("review gate blocks an unfinished revision round", not blocked.ok, blocked.detail)

    (manuscript / "methods.md").write_text("# Methods\n\nConcise revised methods source.\n", encoding="utf-8")
    (manuscript / "discussion.md").write_text("# Discussion\n\nRevised interpretation within the design.\n", encoding="utf-8")
    batch_state = State(batch_project, ".wf").load()
    batch_state.data["current"] = "S24_package_human_review"
    batch_state.stage_info("S24_package_human_review")["status"] = "active"
    batch_state.save()
    mark_outputs = []
    for item_id, rel, summary in [
        ("R001-01", "07_manuscript/methods.md",
         "Condensed the canonical Methods source while retaining reproducibility details."),
        ("R001-02", "07_manuscript/discussion.md",
         "Clarified the canonical Discussion interpretation within the study design."),
    ]:
        mark_outputs.append(subprocess.run(
            [sys.executable, str(ROOT / "tools/rework.py"), "mark", "--item", item_id,
             "--changed-file", rel, "--validated-by", "fixture source validation passed",
             "--summary", summary, "--project", str(batch_project)], capture_output=True,
            text=True, env=batch_env, encoding="utf-8", errors="replace"))
    closed = subprocess.run(
        [sys.executable, str(ROOT / "tools/rework.py"), "close", "--summary",
         "All requested source revisions were completed and validated through the workflow.",
         "--project", str(batch_project)], capture_output=True, text=True, env=batch_env,
        encoding="utf-8", errors="replace")
    batch_state = State(batch_project, ".wf").load()
    closed_gate = revision_gate(Ctx(route_pipe, batch_state, batch_project, review_stage, {}))
    record("completed revision round closes only at its review node with evidence",
           all(proc.returncode == 0 for proc in mark_outputs) and closed.returncode == 0 and
           batch_state.data.get("active_revision_round") is None and closed_gate.ok,
           ((closed.stdout or "") + (closed.stderr or "") + " " + closed_gate.detail).strip()[:180])
    (manuscript / "methods.md").write_text("# Methods\n\nUntracked post-close drift.\n", encoding="utf-8")
    drift_gate = revision_gate(Ctx(route_pipe, batch_state, batch_project, review_stage, {}))
    record("closed-round hash detects untracked source drift", not drift_gate.ok, drift_gate.detail[:150])

    # A later completed round may legitimately supersede an older receipt for the same path.
    import hashlib
    newer = dict(round_payload)
    newer["round_id"] = "R002"
    newer["status"] = "complete"
    newer["completed_at"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    newer["items"] = [{
        "id": "R002-01", "kind": "methods", "status": "done",
        "validated_by": ["fixture second-round validation passed"],
        "changed_files": [{
            "path": "07_manuscript/methods.md",
            "sha256": hashlib.sha256((manuscript / "methods.md").read_bytes()).hexdigest(),
        }],
    }]
    (batch_project / ".wf/revisions/R002.json").write_text(
        json.dumps(newer, indent=2), encoding="utf-8")
    superseded_gate = revision_gate(Ctx(route_pipe, batch_state, batch_project, review_stage, {}))
    record("latest completed round supersedes an older hash for the same source",
           superseded_gate.ok, superseded_gate.detail)
    shutil.rmtree(batch_root, ignore_errors=True)


def run_fulltext_registration(proj: Path) -> None:
    """Legal external full text must be registered and hash-bound before deep reading."""
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    section("full-text registration and deep-read gate")
    refs = proj / "06_refs"
    deep = refs / "deepread"
    deep.mkdir(parents=True, exist_ok=True)
    entries = [
        {"citekey": f"fixture{i}", "pmid": str(90000000 + i),
         "doi": f"10.0000/fixture.{i}", "title": f"Fixture paper {i}",
         "journal": "Fixture Journal", "year": "2025"}
        for i in range(1, 5)
    ]
    (refs / "library.json").write_text(
        json.dumps({"entries": entries}, indent=2), encoding="utf-8")

    env = {**os.environ, "MEDPAPER_PROJECT": str(proj), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    tool = ROOT / "tools/pubmed/fulltext.py"
    selected = []
    registrations_ok = True
    for i, entry in enumerate(entries, 1):
        source = proj / "temp" / f"fixture{i}.pdf"
        source.write_bytes(b"%PDF-1.4\n" + bytes([64 + i]) * 1800)
        p = subprocess.run(
            [sys.executable, str(tool), "register", "--citekey", entry["citekey"],
             "--file", str(source), "--access", "oa",
             "--source-url", f"https://example.org/articles/{i}", "--route", "selftest"],
            capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
        registrations_ok = registrations_ok and p.returncode == 0
        notes = deep / f'{entry["citekey"]}.md'
        notes.write_text(
            f'# {entry["citekey"]}\n\n## Design and population\n' +
            ("Substantive fixture note describing population, methods, estimates, bias, "
             "comparability, limitations and implications for the Discussion. " * 5),
            encoding="utf-8")
        selected.append({
            "citekey": entry["citekey"], "pmid": entry["pmid"],
            "reason": "Needed to benchmark the primary finding against a comparable design.",
            "access": "oa", "fulltext": f'06_refs/fulltext/{entry["citekey"]}.pdf',
            "notes": f'06_refs/deepread/{entry["citekey"]}.md',
        })
    record("four legal local full texts register", registrations_ok)

    manifest_path = refs / "fulltext/retrieval_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    record("retrieval manifest records hashes and sources",
           len(manifest.get("retrievals", [])) == 4 and
           all(len(r.get("sha256", "")) == 64 and
               str(r.get("source_url", "")).startswith("https://")
               for r in manifest.get("retrievals", [])),
           f'{len(manifest.get("retrievals", []))} record(s)')

    (deep / "deepread_index.json").write_text(
        json.dumps({"selected": selected}, indent=2), encoding="utf-8")
    load_all()
    pipe = registry.load()
    state = State(proj, ".wf").load()
    gate = get("deepread_complete")

    def check_deepread():
        return gate(Ctx(pipeline=pipe, state=state, project=proj,
                        stage=pipe.stage("S15_deepread"),
                        spec={"check": "deepread_complete"}))

    outcome = check_deepread()
    record("registered local full texts satisfy deep-read gate", outcome.ok, outcome.detail)

    changed = refs / "fulltext/fixture1.pdf"
    changed.write_bytes(changed.read_bytes() + b"changed")
    outcome = check_deepread()
    record("post-registration file mutation is rejected", not outcome.ok, outcome.detail[:100])

    replacement = proj / "temp/replacement.pdf"
    replacement.write_bytes(b"%PDF-1.4\n" + b"Z" * 1800)
    p = subprocess.run(
        [sys.executable, str(tool), "register", "--citekey", "fixture2",
         "--file", str(replacement), "--access", "oa",
         "--source-url", "https://example.org/replacement"],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("different existing full text is not overwritten implicitly",
           p.returncode != 0 and "already exists" in ((p.stdout or "") + (p.stderr or "")))

    shutil.rmtree(refs, ignore_errors=True)
    for source in (proj / "temp").glob("fixture*.pdf"):
        source.unlink(missing_ok=True)
    replacement.unlink(missing_ok=True)


def run_journal_eligibility(proj: Path) -> None:
    """The journal gate must reject unverified indexing, quartile, and integrity."""
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    section("journal eligibility gate")
    target_dir = proj / "08_submission"
    (target_dir / "cache").mkdir(parents=True, exist_ok=True)
    guidelines_url = "https://journal.example/authors"
    today = date.today().isoformat()
    meta = {
        "journal": "Fixture Clinical Journal", "issn": "1234-5678",
        "publisher": "Fixture Publisher", "scie_indexed": True,
        "scie_source_url": "https://clarivate.example/master-journal-list",
        "scie_verified_at": today, "jcr_quartile": "Q3",
        "jcr_source": "Journal Citation Reports", "jcr_verified_at": today,
        "integrity_status": "clear", "integrity_checked_at": today,
        "integrity_sources": ["https://clarivate.example/journal",
                              "https://publisher.example/journal"],
        "guidelines_url": guidelines_url, "guidelines_fetched_at": today,
        "abbreviation_placement": "central",
        "abbreviation_rule_source": f"Journal is silent; central policy. {guidelines_url}",
        "chosen_by_user": True,
    }
    target = target_dir / "target_journal.json"
    target.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    (target_dir / "guidelines_extract.md").write_text(
        f"# Guidelines\n\nSource: {guidelines_url}\n\n"
        "## Word limits\n3500\n## Reference style\nVancouver\n"
        "## Figure requirements\nTIFF\n## Table requirements\nEditable\n"
        "## Required statements\nEthics and funding\n## Submission items\nDOCX and figures\n",
        encoding="utf-8")
    (target_dir / "cache/guidelines.html").write_text("<html>fixture</html>", encoding="utf-8")

    load_all()
    pipe = registry.load()
    state = State(proj, ".wf").load()
    gate = get("guidelines_sourced")

    def check_journal():
        return gate(Ctx(pipeline=pipe, state=state, project=proj,
                        stage=pipe.stage("S20_journal"),
                        spec={"check": "guidelines_sourced"}))

    outcome = check_journal()
    record("eligible sourced journal passes", outcome.ok, outcome.detail)
    meta["scie_verified_at"] = (date.today() - timedelta(days=91)).isoformat()
    target.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    outcome = check_journal()
    record("stale SCIE verification is rejected", not outcome.ok, outcome.detail[:100])
    meta["scie_verified_at"] = today
    meta["integrity_status"] = "warning"
    target.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    outcome = check_journal()
    record("integrity warning is rejected", not outcome.ok, outcome.detail[:100])
    shutil.rmtree(target_dir, ignore_errors=True)


def run_manuscript_docx(proj: Path) -> None:
    """Canonical assembly and Word normalization must be deterministic and auditable."""
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    section("canonical manuscript and Word package")
    manuscript = proj / "07_manuscript"
    submission = proj / "08_submission"
    bundle = submission / "bundle"
    bundle.mkdir(parents=True, exist_ok=True)
    (manuscript / "title.md").write_text(
        "# Exposure and Clinical Outcome in a Multicentre Cohort Study\n", encoding="utf-8")
    (manuscript / "abstract.md").write_text(
        "# Abstract\n\n## Background\nWe evaluated the association.\n\n"
        "## Methods\nWe analysed 1284 participants.\n\n## Results\nExposure was associated "
        "with outcome (HR 1.87, 95% CI 1.34 to 2.61).\n\n## Conclusions\nThe exposure "
        "was associated with outcome.\n", encoding="utf-8")
    (manuscript / "keywords.md").write_text(
        "Keywords: cohort studies, exposure, prognosis, survival analysis\n", encoding="utf-8")
    (manuscript / "introduction.md").write_text(
        "# Introduction\n\nThe clinical association requires evaluation [@fixture2025].\n",
        encoding="utf-8")
    (manuscript / "methods.md").write_text(
        "# Methods\n\n## Study design and participants\nWe studied 1284 participants.\n\n"
        "## Statistical analysis\nThe primary model estimated an HR with a 95% CI.\n",
        encoding="utf-8")
    (manuscript / "discussion.md").write_text(
        "# Discussion\n\nExposure was associated with outcome (HR 1.87, 95% CI 1.34 to "
        "2.61) (Figure 1).\n", encoding="utf-8")
    (manuscript / "supplementary_methods.md").write_text(
        "# Supplementary Methods\n\nExtended variable coding details are provided here.  \n"
        "Definitions follow without a manual Word line-break control.\n",
        encoding="utf-8")
    env = {**os.environ, "MEDPAPER_PROJECT": str(proj), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    refs = proj / "06_refs"
    refs.mkdir(parents=True, exist_ok=True)
    bib = refs / "refs.bib"
    bib.write_text("@article{fixture2025, title={Fixture cohort report}, "
                   "author={Author, Alice}, journal={Fixture Journal}, year={2025}}\n",
                   encoding="utf-8")
    (refs / "refs.ris").write_text(
        "TY  - JOUR\nID  - fixture2025\nTI  - Fixture cohort report\nPY  - 2025\nER  - \n",
        encoding="utf-8")
    assembler = ROOT / "tools/manuscript/assemble.py"
    proc = subprocess.run([sys.executable, str(assembler)], capture_output=True, text=True,
                          env=env, encoding="utf-8", errors="replace")
    record("full manuscript assembles", proc.returncode == 0,
           ((proc.stdout or "") + (proc.stderr or "")).strip()[:100])

    load_all()
    pipe = registry.load()
    state = State(proj, ".wf").load()

    def gate(check_name: str, stage: str, **spec):
        return get(check_name)(Ctx(pipeline=pipe, state=state, project=proj,
                                   stage=pipe.stage(stage),
                                   spec={"check": check_name, **spec}))

    # S19 must produce a current, versioned package that can be handed to a third party.
    review_report = manuscript / "independent_publishability_review.md"
    review_report.write_text(
        "# Independent publishability review\n\n## Verdict\nREADY.\n\n"
        "## Critical barriers\nNone.\n\n## Scientific validity\nAcceptable.\n\n"
        "## Reporting completeness\nComplete.\n\n"
        "## Tables and supplementary material\nConsistent.\n\n"
        "## Journal suitability\nA realistic SCIE journal is plausible.\n\n"
        "## Required revisions\nNone before user review.\n\n"
        "## Post-revision outlook\nProceed after explicit user review.\n",
        encoding="utf-8",
    )
    prior_state = json.loads(json.dumps(state.data))
    state.data["current"] = "S19_human_review"
    state.stage_info("S19_human_review")["status"] = "active"
    state.save()
    review_tool = ROOT / "tools/manuscript/review_package.py"
    first_review_zip = subprocess.run(
        [sys.executable, str(review_tool), "build", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    review_manifest_path = manuscript / "review_packages/latest_review_package.json"
    review_manifest = (json.loads(review_manifest_path.read_text(encoding="utf-8"))
                       if review_manifest_path.is_file() else {})
    review_archive = proj / review_manifest.get("archive_path", "missing")
    members = set()
    if review_archive.is_file():
        with zipfile.ZipFile(review_archive) as zf:
            members = set(zf.namelist())
    outcome = gate("manuscript_review_package_current", "S19_human_review")
    record("S19 builds a verified third-party review ZIP with exact review materials",
           first_review_zip.returncode == 0 and outcome.ok and
           review_manifest.get("package_revision") == 1 and
           {"07_manuscript/full_manuscript.md",
            "07_manuscript/supplementary_methods.md",
            "07_manuscript/independent_publishability_review.md",
            "06_refs/refs.bib", "06_refs/refs.ris",
            "REVIEW_README.txt", "review_manifest.json"}.issubset(members) and
           any(name.startswith("04_tables/main/") for name in members) and
           any(name.startswith("04_tables/supplementary/") for name in members) and
           any(name.startswith("05_figures/out/") and name.endswith(".png") for name in members) and
           not any(name.startswith(("02_data/", "03_analysis/", "06_refs/fulltext/"))
                   for name in members),
           ((first_review_zip.stdout or "") + (first_review_zip.stderr or "") +
            outcome.detail).strip()[:180])
    unchanged_review_zip = subprocess.run(
        [sys.executable, str(review_tool), "build", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    unchanged_manifest = json.loads(review_manifest_path.read_text(encoding="utf-8"))
    record("unchanged S19 sources reuse the verified ZIP without duplicate versions",
           unchanged_review_zip.returncode == 0 and "reused unchanged" in unchanged_review_zip.stdout and
           unchanged_manifest.get("archive_path") == review_manifest.get("archive_path") and
           unchanged_manifest.get("package_revision") == 1,
           unchanged_review_zip.stdout.strip()[:140])

    state.record_decision(
        "manuscript_human_reviewed", "YES",
        "The fixture intentionally records an old ambiguous approval value for rejection.")
    wrong_approval = gate("s19_review_release_explicit", "S19_human_review")
    state.record_decision(
        "manuscript_human_reviewed", "NO_FURTHER_REVIEW",
        "The user explicitly stated that no further scientific review was needed after ZIP "
        f"v001 package {review_manifest['package_id'][:12]}.")
    explicit_approval = gate("s19_review_release_explicit", "S19_human_review")
    record("S19 rejects generic approval and accepts only explicit no-further-review state",
           not wrong_approval.ok and explicit_approval.ok,
           f"old={wrong_approval.detail}; explicit={explicit_approval.detail}")

    review_report.write_text(
        review_report.read_text(encoding="utf-8") +
        "\nThe reviewer clarified one editorial point for the revised version.\n",
        encoding="utf-8",
    )
    stale_review = gate("manuscript_review_package_current", "S19_human_review")
    rebuilt_review_zip = subprocess.run(
        [sys.executable, str(review_tool), "build", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    revised_manifest = json.loads(review_manifest_path.read_text(encoding="utf-8"))
    revised_review = gate("manuscript_review_package_current", "S19_human_review")
    stale_approval = gate("s19_review_release_explicit", "S19_human_review")
    record("changed S19 review material invalidates v001 and creates verified v002",
           not stale_review.ok and rebuilt_review_zip.returncode == 0 and revised_review.ok and
           not stale_approval.ok and
           revised_manifest.get("package_revision") == 2 and
           revised_manifest.get("archive_path") != review_manifest.get("archive_path") and
           review_archive.is_file() and (proj / revised_manifest["archive_path"]).is_file(),
           ((rebuilt_review_zip.stdout or "") + revised_review.detail).strip()[:180])

    # Explicit approval creates an immutable scientific master; journal work is a derived copy.
    (manuscript / "human_review.md").write_text(
        "# Human review\n\n## Materials presented\nS19 v002.\n\n"
        "## Independent verdict\nREADY.\n\n## User-requested revisions\nNone.\n\n"
        "## Revalidation\nPassed.\n\n## Approval to proceed\nNo further review.\n",
        encoding="utf-8",
    )
    state.record_decision(
        "manuscript_human_reviewed", "NO_FURTHER_REVIEW",
        "The user explicitly stated no further review was needed for current ZIP v002 "
        f"package {revised_manifest['package_id'][:12]}.")
    scientific_tool = ROOT / "tools/manuscript/scientific_freeze.py"
    frozen_science = subprocess.run(
        [sys.executable, str(scientific_tool), "freeze", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    frozen_gate = gate("scientific_master_frozen", "S19_human_review")
    record("S19 freezes the exact user-approved journal-independent scientific master",
           frozen_science.returncode == 0 and frozen_gate.ok,
           ((frozen_science.stdout or "") + (frozen_science.stderr or "") +
            frozen_gate.detail).strip()[:180])
    approved_full = (manuscript / "full_manuscript.md").read_text(encoding="utf-8")
    (manuscript / "full_manuscript.md").write_text(
        approved_full + "\nDetached post-freeze edit.\n", encoding="utf-8")
    changed_science = subprocess.run(
        [sys.executable, str(scientific_tool), "verify", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("post-S19 scientific-master drift is rejected",
           changed_science.returncode != 0 and
           any(marker in (changed_science.stderr or "") for marker in
               ("scientific source changed after freeze", "review materials changed")),
           (changed_science.stderr or "").strip()[:150])
    (manuscript / "full_manuscript.md").write_text(approved_full, encoding="utf-8")

    (submission / "target_journal.json").write_text(json.dumps({
        "journal": "Fixture Clinical Journal", "issn": "1234-5678",
        "chosen_by_user": True, "guidelines_url": "https://journal.example/authors",
        "guidelines_fetched_at": date.today().isoformat(),
    }, indent=2), encoding="utf-8")
    workspace_tool = ROOT / "tools/manuscript/journal_workspace.py"
    workspace_init = subprocess.run(
        [sys.executable, str(workspace_tool), "init", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    workspace_gate = gate("journal_workspace_ready", "S20_journal", require_pristine=True)
    record("S20 derives a pristine journal integration layer from the frozen master",
           workspace_init.returncode == 0 and workspace_gate.ok and
           (submission / "integration/full_manuscript.md").is_file(),
           ((workspace_init.stdout or "") + (workspace_init.stderr or "") +
            workspace_gate.detail).strip()[:180])
    integration_full = submission / "integration/full_manuscript.md"
    integration_full.write_text(
        integration_full.read_text(encoding="utf-8").replace("evaluated", "assessed", 1),
        encoding="utf-8")
    unchanged_science = subprocess.run(
        [sys.executable, str(scientific_tool), "verify", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("journal-specific integration edits do not alter the frozen scientific master",
           unchanged_science.returncode == 0,
           ((unchanged_science.stdout or "") + (unchanged_science.stderr or "")).strip()[:140])
    state.data = prior_state
    state.save()

    outcome = gate("abbreviations_centralized", "S21_authors")
    record("short display abbreviation lists may remain local", outcome.ok, outcome.detail)
    captions = proj / "04_tables/table_captions.md"
    captions_good = captions.read_text(encoding="utf-8")
    captions.write_text(captions_good +
                        "\nAbbreviations: BMI, body mass index; eGFR, estimated glomerular "
                        "filtration rate; HbA1c, glycated haemoglobin; SBP, systolic blood pressure; "
                        "DBP, diastolic blood pressure; CKD, chronic kidney disease.\n",
                        encoding="utf-8")
    outcome = gate("abbreviations_centralized", "S21_authors")
    record("crowded display abbreviations require centralization", not outcome.ok,
           outcome.detail[:100])
    captions.write_text(captions_good, encoding="utf-8")

    # A realistic centralized glossary must pass without redefining panel/table labels.
    scoped = proj / "temp/abbreviations_case"
    (scoped / "05_figures").mkdir(parents=True)
    (scoped / "04_tables").mkdir(parents=True)
    (scoped / "07_manuscript").mkdir(parents=True)
    (scoped / "05_figures/legends.md").write_text(
        "# Figure legends\n\n## Figure 1.\nModels for eGFR and HbA1c. "
        "Abbreviations are listed in Declarations and Statements.\n", encoding="utf-8")
    (scoped / "04_tables/table_captions.md").write_text(
        "# Table captions\n\n## Table S1.\nClinical characteristics. "
        "Abbreviations are listed in Declarations and Statements.\n", encoding="utf-8")
    central = ("# Declarations and Statements\n\n## Abbreviations\n\n"
               "eGFR, estimated glomerular filtration rate; HbA1c, glycated haemoglobin.\n")
    (scoped / "07_manuscript/statements.md").write_text(central, encoding="utf-8")
    scoped_full = scoped / "07_manuscript/full_manuscript.md"
    scoped_full.write_text("# Discussion\n\nClinical interpretation.\n\n" + central +
                           "\n# References\n", encoding="utf-8")
    ctx = Ctx(pipeline=pipe, state=state, project=scoped, stage=pipe.stage("S17_assemble"),
              spec={"check": "abbreviations_centralized"})
    outcome = get("abbreviations_centralized")(ctx)
    record("central glossary accepts mixed-case terms and table labels", outcome.ok, outcome.detail)
    scoped_full.write_text("# Discussion\n\nClinical interpretation.\n", encoding="utf-8")
    outcome = get("abbreviations_centralized")(ctx)
    record("central glossary cannot be omitted from full manuscript", not outcome.ok,
           outcome.detail[:100])
    shutil.rmtree(scoped)

    outcome = gate("manuscript_structure", "S17_assemble")
    record("canonical order and keywords pass", outcome.ok, outcome.detail)
    outcome = gate("assembly_matches_sources", "S17_assemble")
    record("canonical manuscript exactly matches source sections", outcome.ok, outcome.detail)
    full = manuscript / "full_manuscript.md"
    good = full.read_text(encoding="utf-8")
    full.write_text(good.replace("cohort studies, exposure", "cohort studies; exposure"),
                    encoding="utf-8")
    outcome = gate("manuscript_structure", "S17_assemble")
    record("malformed keyword separators are rejected", not outcome.ok, outcome.detail[:100])
    full.write_text(good, encoding="utf-8")

    # S17 can assemble abbreviation definitions before author/admin collection.
    statement_source = manuscript / "statements.md"
    statement_source.write_text(central, encoding="utf-8")
    assembled = subprocess.run([sys.executable, str(assembler)], capture_output=True, text=True,
                               env=env, encoding="utf-8", errors="replace")
    with_glossary = full.read_text(encoding="utf-8")
    record("assembly includes glossary before References without author intake",
           assembled.returncode == 0 and
           with_glossary.index("# Discussion") < with_glossary.index("# Declarations and Statements") <
           with_glossary.index("# References"))
    statement_source.unlink()
    full.write_text(good, encoding="utf-8")
    full.write_text(good.replace("clinical association", "clinical relationship", 1),
                    encoding="utf-8")
    outcome = gate("assembly_matches_sources", "S17_assemble")
    record("source-section omission or divergence is rejected", not outcome.ok,
           outcome.detail[:100])
    outcome = gate("assembly_matches_sources", "S19_human_review")
    record("S19 rejects a detached full-manuscript edit", not outcome.ok,
           outcome.detail[:100])
    full.write_text("# Exposure and Clinical Outcome in a Multicentre Cohort Study\n\n" +
                    good.split("\n\n", 1)[1].split("# Figure legends", 1)[0] +
                    "# Figure legends\n\n", encoding="utf-8")
    outcome = gate("manuscript_structure", "S17_assemble")
    record("empty planned figure legend is rejected", not outcome.ok, outcome.detail[:100])
    full.write_text(good, encoding="utf-8")

    # A title-page count is optional, but when present it must be derived from the
    # distinct citekeys used by the canonical manuscript rather than library size.
    integration = submission / "integration"
    integration_full = integration / "full_manuscript.md"
    title_page = integration / "title_page.md"
    title_page.write_text(
        "# Exposure and Clinical Outcome in a Multicentre Cohort Study\n\n"
        "Fixture Author\n\nNumber of references: 50\n", encoding="utf-8")
    outcome = gate("title_page_reference_count", "S21_authors",
                   manuscript="08_submission/integration/full_manuscript.md",
                   title_page="08_submission/integration/title_page.md")
    record("incorrect title-page reference count is rejected", not outcome.ok,
           outcome.detail[:110])
    counter = ROOT / "tools/manuscript/reference_count.py"
    synced = subprocess.run(
        [sys.executable, str(counter), "sync", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("actual citation count synchronizes title page",
           synced.returncode == 0 and "Number of references: 1" in
           title_page.read_text(encoding="utf-8"),
           ((synced.stdout or "") + (synced.stderr or "")).strip()[:110])
    outcome = gate("title_page_reference_count", "S21_authors",
                   manuscript="08_submission/integration/full_manuscript.md",
                   title_page="08_submission/integration/title_page.md")
    record("title-page reference count gate passes after synchronization", outcome.ok,
           outcome.detail)

    title_page.write_text(
        "# Exposure and Clinical Outcome in a Multicentre Cohort Study\n\nFixture Author\n",
        encoding="utf-8")
    outcome = gate("title_page_reference_count", "S21_authors",
                   manuscript="08_submission/integration/full_manuscript.md",
                   title_page="08_submission/integration/title_page.md")
    record("reference count remains optional when journal does not request it", outcome.ok,
           outcome.detail)
    added = subprocess.run(
        [sys.executable, str(counter), "sync", "--add", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("reference count field is added only on explicit request",
           added.returncode == 0 and title_page.read_text(encoding="utf-8").count(
               "Number of references: 1") == 1,
           ((added.stdout or "") + (added.stderr or "")).strip()[:110])

    style = {
        "font_family": "Times New Roman", "body_font_pt": 12, "title_font_pt": 14,
        "section_heading_font_pt": 12, "subsection_heading_font_pt": 12,
        "line_spacing": 2.0, "paragraph_spacing_after_pt": 0, "margins_in": 1.0,
        "paper_size": "A4", "all_text_black": True, "external_hyperlinks": False,
        "source": "Fixture journal is silent; documented medical-manuscript fallback. "
                  "https://journal.example/authors",
        "fallbacks": ["font_family", "body_font_pt", "title_font_pt",
                      "section_heading_font_pt", "subsection_heading_font_pt",
                      "line_spacing", "paragraph_spacing_after_pt", "margins_in",
                      "paper_size"],
    }
    style_path = submission / "docx_style.json"
    style_path.write_text(json.dumps(style, indent=2), encoding="utf-8")
    outcome = gate("docx_style_config", "S20_journal")
    record("journal Word style configuration passes", outcome.ok, outcome.detail)
    bad_style = dict(style)
    bad_style["font_family"] = "Arial"
    style_path.write_text(json.dumps(bad_style, indent=2), encoding="utf-8")
    outcome = gate("docx_style_config", "S20_journal")
    record("journal-silent non-Times font is rejected", not outcome.ok, outcome.detail[:100])
    style_path.write_text(json.dumps(style, indent=2), encoding="utf-8")

    cover = submission / "cover_letter.md"
    cover.write_text("# Cover letter\n\nDear Editor,\n\nPlease consider this cohort study.\n",
                     encoding="utf-8")
    builder = ROOT / "tools/manuscript/build_docx.py"

    def build(kind: str, source: Path, name: str) -> bool:
        command = [sys.executable, str(builder), "build", "--kind", kind,
                   "--input", str(source), "--output", str(bundle / name),
                   "--style-config", str(style_path)]
        if kind == "manuscript":
            command += ["--bibliography", str(bib)]
        proc_ = subprocess.run(
            command, capture_output=True, text=True,
            env=env, encoding="utf-8", errors="replace")
        record(f"{kind} DOCX builds and self-audits", proc_.returncode == 0,
               ((proc_.stdout or "") + (proc_.stderr or "")).strip()[:120])
        return proc_.returncode == 0

    builds_ok = all([
        build("manuscript", integration_full, "manuscript.docx"),
        build("title_page", title_page, "title_page.docx"),
        build("cover_letter", cover, "cover_letter.docx"),
        build("supplementary", integration / "supplementary_methods.md", "supplementary_methods.docx"),
    ])
    word_xmls = []
    for name in ("manuscript.docx", "title_page.docx", "cover_letter.docx",
                 "supplementary_methods.docx"):
        with zipfile.ZipFile(bundle / name) as zf:
            word_xmls.extend(zf.read(item) for item in zf.namelist()
                             if item.startswith("word/") and item.endswith(".xml"))
    forbidden_controls = (b"<w:outlineLvl", b"<w:keepNext", b"<w:keepLines",
                          b"<w:pageBreakBefore", b"<w:pBdr")
    record("all Word XML parts remove black-square and folding controls",
           builds_ok and not any(tag in blob for tag in forbidden_controls for blob in word_xmls))
    with zipfile.ZipFile(bundle / "manuscript.docx") as zf:
        manuscript_xml = zf.read("word/document.xml")
    record("visible section headings are flattened to SectionHeading",
           b'w:val="SectionHeading"' in manuscript_xml and
           b'w:val="Heading1"' not in manuscript_xml and b'w:val="Heading2"' not in manuscript_xml)
    manuscript_text = "\n".join(
        paragraph.text for paragraph in __import__("docx").Document(bundle / "manuscript.docx").paragraphs)
    record("Figure legends section title appears exactly once in Word",
           sum(line.strip().casefold() == "figure legends"
               for line in manuscript_text.splitlines()) == 1)

    supplement_source = manuscript / "supplementary_methods.md"
    supplement_good = supplement_source.read_text(encoding="utf-8")
    supplement_source.write_text(
        "# Supplementary Methods\n\n| Variable | Definition |\n|---|---|\n| A | B |\n\n---\n",
        encoding="utf-8")
    outcome = gate("supplementary_methods_clean", "S08_methods")
    record("supplementary Methods rejects embedded tables and Markdown rules",
           not outcome.ok, outcome.detail[:120])
    invalid_supplement = bundle / "invalid_supplement.docx"
    invalid_proc = subprocess.run(
        [sys.executable, str(builder), "build", "--kind", "supplementary",
         "--input", str(supplement_source), "--output", str(invalid_supplement),
         "--style-config", str(style_path)], capture_output=True, text=True, env=env,
        encoding="utf-8", errors="replace")
    record("Word builder refuses a supplementary-Methods table bypass",
           invalid_proc.returncode != 0 and not invalid_supplement.exists(),
           ((invalid_proc.stdout or "") + (invalid_proc.stderr or "")).strip()[:120])
    supplement_source.write_text(supplement_good, encoding="utf-8")

    duplicate_source = submission / "duplicate_legend_source.md"
    duplicate_source.write_text(full.read_text(encoding="utf-8") +
                                "\n# Figure legends\n", encoding="utf-8")
    duplicate_output = bundle / "duplicate_legend.docx"
    duplicate_proc = subprocess.run(
        [sys.executable, str(builder), "build", "--kind", "manuscript",
         "--input", str(duplicate_source), "--output", str(duplicate_output),
         "--style-config", str(style_path), "--bibliography", str(bib)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("Word builder refuses duplicate Figure legends titles",
           duplicate_proc.returncode != 0 and not duplicate_output.exists(),
           ((duplicate_proc.stdout or "") + (duplicate_proc.stderr or "")).strip()[:120])
    duplicate_source.unlink(missing_ok=True)

    with zipfile.ZipFile(bundle / "supplementary_methods.docx") as zf:
        supplement_xml = zf.read("word/document.xml")
    record("manual Word line-break controls are removed",
           b"<w:br" not in supplement_xml and b"<w:cr" not in supplement_xml)
    from docx import Document
    from docx.enum.text import WD_BREAK
    from manuscript.build_docx import audit_docx, normalize_docx
    tampered_supplement = bundle / "tampered_supplement.docx"
    shutil.copy2(bundle / "supplementary_methods.docx", tampered_supplement)
    tampered_doc = Document(tampered_supplement)
    tampered_doc.add_table(rows=2, cols=2)
    tampered_doc.save(tampered_supplement)
    record("Word audit rejects an embedded supplementary-Methods table",
           any("supplementary Methods contains a table" in issue for issue in
               audit_docx(tampered_supplement, style, "supplementary")))
    tampered_supplement.unlink()
    raw_breaks = bundle / "line_break_fixture.docx"
    raw_doc = Document()
    paragraph = raw_doc.add_paragraph("First address line")
    paragraph.add_run().add_break()
    paragraph.add_run("Second address line").italic = True
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("After page break")
    cell = raw_doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "First cell line\nSecond cell line"
    raw_doc.save(raw_breaks)
    normalize_docx(raw_breaks, style, "cover_letter")
    normalized = Document(raw_breaks)
    record("line-break conversion preserves paragraphs, italics and table lines",
           normalized.paragraphs[0].text == "First address line" and
           normalized.paragraphs[1].text.startswith("Second address line") and
           any(run.italic for run in normalized.paragraphs[1].runs) and
           [p.text for p in normalized.tables[0].cell(0, 0).paragraphs] ==
           ["First cell line", "Second cell line"])
    with zipfile.ZipFile(raw_breaks) as zf:
        record("explicit page breaks survive manual-break cleanup",
               b'w:type="page"' in zf.read("word/document.xml"))
    raw_breaks.unlink()

    arrow_source = submission / "arrow_source.md"
    arrow_source.write_text("# Cover letter\n\nA downward marker \u2193 must not survive.\n",
                            encoding="utf-8")
    arrow_output = bundle / "arrow_output.docx"
    arrow_proc = subprocess.run(
        [sys.executable, str(builder), "build", "--kind", "cover_letter",
         "--input", str(arrow_source), "--output", str(arrow_output),
         "--style-config", str(style_path)], capture_output=True, text=True, env=env,
        encoding="utf-8", errors="replace")
    record("Word builder rejects a literal down arrow",
           arrow_proc.returncode != 0 and not arrow_output.exists(),
           ((arrow_proc.stdout or "") + (arrow_proc.stderr or "")).strip()[:100])
    arrow_source.unlink(missing_ok=True)
    (bundle / "SUBMISSION_CHECKLIST.md").write_text("# Submission checklist\n\n- Complete\n",
                                                     encoding="utf-8")
    shutil.copy2(proj / "05_figures/out/Figure1.tiff", bundle / "Figure1.tiff")
    shutil.copy2(proj / "04_tables/main/Table1.xlsx", bundle / "Table1.xlsx")
    items = [
        ("manuscript", "manuscript.docx"), ("title_page", "title_page.docx"),
        ("cover_letter", "cover_letter.docx"),
        ("supplementary", "supplementary_methods.docx"),
        ("figures", "Figure1.tiff"), ("tables", "Table1.xlsx"),
        ("checklist", "SUBMISSION_CHECKLIST.md"),
    ]
    manifest = {"built_at": "2026-09-04", "journal": "Fixture Clinical Journal",
                "items": [{"role": role, "file": f"08_submission/bundle/{name}",
                           "required_by": "guidelines_extract.md > Submission items"}
                          for role, name in items]}
    manifest_path = bundle / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    outcome = gate("docx_bundle_ready", "S23_package")
    record("Word bundle gate accepts normalized DOCX files", builds_ok and outcome.ok,
           outcome.detail)
    outcome = gate("title_page_reference_count", "S23_package",
                   manuscript="08_submission/integration/full_manuscript.md",
                   title_page="08_submission/integration/title_page.md")
    record("Word title-page reference count matches canonical manuscript", outcome.ok,
           outcome.detail)
    title_page_docx = bundle / "title_page.docx"
    approved_title_page = title_page_docx.read_bytes()
    altered_title = Document(title_page_docx)
    for paragraph in altered_title.paragraphs:
        if "Number of references:" in paragraph.text:
            paragraph.text = "Number of references: 50"
    altered_title.save(title_page_docx)
    outcome = gate("title_page_reference_count", "S23_package",
                   manuscript="08_submission/integration/full_manuscript.md",
                   title_page="08_submission/integration/title_page.md")
    record("manually altered Word reference count is rejected", not outcome.ok,
           outcome.detail[:110])
    title_page_docx.write_bytes(approved_title_page)

    from wfcore.packagecontent import write_baseline
    baseline_path = write_baseline(proj)
    outcome = gate("package_content_matches_baseline", "S23_package")
    record("S23 captures a valid submission visible-text baseline",
           baseline_path.is_file() and outcome.ok, outcome.detail)
    baseline_cli = ROOT / "tools/package_content.py"
    blocked_capture = subprocess.run(
        [sys.executable, str(baseline_cli), "capture", "--replace", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("content baseline cannot be recaptured outside S23",
           blocked_capture.returncode != 0 and "allowed only at S23_package" in
           (blocked_capture.stderr or ""), (blocked_capture.stderr or "").strip()[:110])

    format_only_doc = Document(title_page_docx)
    format_only_doc.paragraphs[1].runs[0].bold = True
    format_only_doc.save(title_page_docx)
    outcome = gate("package_content_matches_baseline", "S24_package_human_review")
    record("Word formatting-only edits preserve the content baseline",
           outcome.ok and "formatting-only" in outcome.detail, outcome.detail)
    title_page_docx.write_bytes(approved_title_page)

    text_changed_doc = Document(title_page_docx)
    text_changed_doc.paragraphs[1].text = "Different Author"
    text_changed_doc.save(title_page_docx)
    outcome = gate("package_content_matches_baseline", "S24_package_human_review")
    record("direct Word text edits are rejected", not outcome.ok, outcome.detail[:110])
    title_page_docx.write_bytes(approved_title_page)

    approved_cover_source = cover.read_text(encoding="utf-8")
    cover.write_text(approved_cover_source + "\nA detached source edit.\n", encoding="utf-8")
    outcome = gate("package_content_matches_baseline", "S24_package_human_review")
    record("source edits without an S23 rebuild are rejected", not outcome.ok,
           outcome.detail[:110])
    cover.write_text(approved_cover_source, encoding="utf-8")
    outcome = gate("package_content_matches_baseline", "S24_package_human_review")
    record("restoring Word text and sources restores the content baseline", outcome.ok,
           outcome.detail)

    proc = subprocess.run(
        [sys.executable, str(builder), "audit", "--manifest", str(manifest_path),
         "--style-config", str(style_path)], capture_output=True, text=True, env=env,
        encoding="utf-8", errors="replace")
    record("standalone Word manifest audit passes", proc.returncode == 0,
           ((proc.stdout or "") + (proc.stderr or "")).strip()[:120])

    # Freeze the exact package the user approved and reject any post-confirmation mutation.
    today = date.today().isoformat()
    guidelines_url = "https://journal.example/authors"
    target_payload = json.loads((submission / "target_journal.json").read_text(encoding="utf-8"))
    target_payload.update({"guidelines_url": guidelines_url, "guidelines_fetched_at": today})
    (submission / "target_journal.json").write_text(
        json.dumps(target_payload, indent=2), encoding="utf-8")
    (submission / "guidelines_extract.md").write_text(
        f"# Guidelines\n\nSource: {guidelines_url}\n\n## Submission items\nAll fixture files.\n",
        encoding="utf-8")
    (submission / "submission_qc.md").write_text(
        "# Submission QC\n\n## Rendered files inspected\nAll.\n", encoding="utf-8")
    (submission / "cache").mkdir(exist_ok=True)
    (submission / "cache/guidelines.html").write_text("<html>fixture</html>", encoding="utf-8")
    freezer = ROOT / "tools/package_review.py"
    frozen = subprocess.run(
        [sys.executable, str(freezer), "freeze", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    outcome = gate("bundle_matches_freeze", "S24_package_human_review")
    record("user-approved submission package freezes and verifies",
           frozen.returncode == 0 and outcome.ok,
           ((frozen.stdout or "") + (frozen.stderr or "") + outcome.detail).strip()[:140])
    manuscript_docx = bundle / "manuscript.docx"
    approved_bytes = manuscript_docx.read_bytes()
    manuscript_docx.write_bytes(approved_bytes + b"post-confirmation-change")
    changed = subprocess.run(
        [sys.executable, str(freezer), "verify", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    outcome = gate("bundle_matches_freeze", "S25_submission_audit")
    record("post-confirmation package mutation is rejected",
           changed.returncode != 0 and not outcome.ok,
           ((changed.stderr or "") + outcome.detail).strip()[:140])
    manuscript_docx.write_bytes(approved_bytes)
    verified = subprocess.run(
        [sys.executable, str(freezer), "verify", "--project", str(proj)],
        capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("restoring the approved package clears freeze gate", verified.returncode == 0,
           ((verified.stdout or "") + (verified.stderr or "")).strip()[:120])
    freeze_payload = json.loads((submission / "package_review_freeze.json").read_text(encoding="utf-8"))
    record("final freeze includes the S23 content baseline",
           any(item.get("path") == "08_submission/package_content_baseline.json"
               for item in freeze_payload.get("files", [])))
    audit_path = submission / "independent_submission_audit.md"
    audit_path.write_text("# Verdict\n\nPASS\n", encoding="utf-8")
    outcome = gate("submission_audit_matches_freeze", "S25_submission_audit")
    record("stale or unbound independent audit is rejected", not outcome.ok, outcome.detail[:120])
    audit_path.write_text(
        "# Verdict\n\nPASS\n\nFreeze ID: " + freeze_payload["freeze_id"] +
        "\n\n# Guideline compliance\n\nPass.\n\n# Required files and omissions\n\nNone.\n"
        "\n# Cross-file consistency\n\nPass.\n\n# Formatting and technical checks\n\nPass.\n"
        "\n# Issues requiring correction\n\nNone.\n\n# Final recommendation\n\nReady.\n",
        encoding="utf-8")
    outcome = gate("submission_audit_matches_freeze", "S25_submission_audit")
    record("independent audit binds to the exact approved freeze", outcome.ok, outcome.detail)

    missing_supp = dict(manifest)
    missing_supp["items"] = [item for item in manifest["items"]
                              if item["role"] != "supplementary"]
    manifest_path.write_text(json.dumps(missing_supp, indent=2), encoding="utf-8")
    outcome = gate("docx_bundle_ready", "S23_package")
    record("missing supplementary Methods DOCX is rejected", not outcome.ok,
           outcome.detail[:100])

    def make_blue_docx(source: Path, destination: Path) -> bool:
        with zipfile.ZipFile(source, "r") as src:
            entries = [(info, src.read(info.filename)) for info in src.infolist()]
        changed = False
        with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as out:
            for info, blob in entries:
                if info.filename == "word/document.xml":
                    replaced = blob.replace(b'w:val="000000"', b'w:val="0000FF"', 1)
                    changed = replaced != blob
                    blob = replaced
                out.writestr(info, blob)
        return changed

    blue = bundle / "manuscript_blue.docx"
    changed = make_blue_docx(bundle / "manuscript.docx", blue)
    bad_manifest = json.loads(json.dumps(manifest))
    next(item for item in bad_manifest["items"] if item["role"] == "manuscript")["file"] = (
        "08_submission/bundle/manuscript_blue.docx")
    manifest_path.write_text(json.dumps(bad_manifest, indent=2), encoding="utf-8")
    outcome = gate("docx_bundle_ready", "S23_package")
    record("blue Word text is rejected", changed and not outcome.ok, outcome.detail[:100])
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    blue.unlink(missing_ok=True)

    if os.environ.get("MEDPAPER_SELFTEST_KEEP_ARTIFACTS") == "1":
        qa_docs = proj.parent / "word_qa"
        qa_docs.mkdir(parents=True, exist_ok=True)
        for name in ("manuscript.docx", "supplementary_methods.docx"):
            shutil.copy2(bundle / name, qa_docs / name)
    shutil.rmtree(submission, ignore_errors=True)
    for name in ("title.md", "abstract.md", "keywords.md", "introduction.md", "methods.md",
                 "discussion.md", "supplementary_methods.md", "full_manuscript.md", "title_page.md"):
        (manuscript / name).unlink(missing_ok=True)
    bib.unlink(missing_ok=True)

def run_polish(proj: Path) -> None:
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    load_all()
    pipe = registry.load()
    st = State(proj, ".wf").load()
    tool = ROOT / "tools/text/polish.py"
    env = {**os.environ, "MEDPAPER_PROJECT": str(proj), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}

    def polish(*argv):
        return subprocess.run([sys.executable, str(tool), *argv], capture_output=True,
                              text=True, env=env, encoding="utf-8", errors="replace")

    def gate(name: str, stage: str = "S22_polish", **spec):
        fn = get(name)
        return fn(Ctx(pipeline=pipe, state=st, project=proj,
                      stage=pipe.stage(stage), spec={"check": name, **spec}))

    integration = proj / "08_submission/integration"
    integration.mkdir(parents=True, exist_ok=True)
    disc = integration / "full_manuscript.md"

    section("polish: linting AI slop")
    disc.write_text(SLOP, encoding="utf-8")
    p = polish("snapshot")
    record("snapshot taken", p.returncode == 0 and
           (integration / "prepolish/facts.json").exists())
    record("snapshot refuses to overwrite silently", polish("snapshot").returncode == 1)

    p = polish("lint")
    out = p.stdout or ""
    counts = json.loads((integration / "polish_report.json").read_text(encoding="utf-8"))["counts"]
    record("tier-A AI phrases detected", counts["ai_tier_a"] >= 10, f"{counts['ai_tier_a']} found")
    record("structural tells detected", counts["structure_blocking"] >= 1,
           f"{counts['structure_blocking']} blocking")
    record("house-style defects detected", counts["style_blocking"] >= 4,
           f"{counts['style_blocking']} blocking")
    for want, label in [("utilize", "inflated verb"), ("range_dash", "hyphen numeric range"),
                        ("unit_spacing", "value glued to unit"), ("data_agreement", "'data was'"),
                        ("p_zero", "P = 0.000"), ("causal_overclaim", "causal overclaim"),
                        ("hedge_stacking", "stacked hedges")]:
        record(f"caught: {label}", want in out, "" if want in out else f"{want!r} absent from report")

    r = gate("ai_tells_clean")
    record("ai_tells_clean rejects the slop", not r.ok, r.detail[:100])
    r = gate("style_consistent")
    record("style_consistent rejects the slop", not r.ok, r.detail[:100])

    section("polish: fact preservation")
    disc.write_text(CLEAN, encoding="utf-8")
    p = polish("diff")
    record("polished text preserves every fact", p.returncode == 0,
           (p.stdout or "").strip().splitlines()[-1][:90])
    polish("lint")
    r = gate("ai_tells_clean")
    record("ai_tells_clean accepts the polished text", r.ok, r.detail[:100])
    r = gate("style_consistent")
    record("style_consistent accepts the polished text", r.ok, r.detail[:100])

    disc.write_text(CLEAN.replace("HR 1.87, 95% CI\n1.34 to 2.61", "HR 1.90, 95% CI\n1.40 to 2.60"),
                    encoding="utf-8")
    p = polish("diff")
    record("altered statistic is caught by diff", p.returncode == 2,
           next((ln.strip() for ln in (p.stdout or "").splitlines() if "LOST" in ln), "")[:90])
    r = gate("polish_preserves_facts")
    record("polish_preserves_facts gate rejects it", not r.ok, r.detail[:100])

    disc.write_text(CLEAN.replace("[@smith2020cohort]", ""), encoding="utf-8")
    record("dropped citation is caught by diff", polish("diff").returncode == 2)

    disc.write_text(CLEAN.replace("(Figure 1A)", ""), encoding="utf-8")
    record("dropped figure reference is caught by diff", polish("diff").returncode == 2)

    disc.write_text(CLEAN, encoding="utf-8")
    record("restoring the text clears diff", polish("diff").returncode == 0)

    section("polish: journal limits")
    (proj / "08_submission").mkdir(parents=True, exist_ok=True)
    gx = proj / "08_submission/guidelines_extract.md"
    gx.write_text("# Guidelines\n\n## Word limits\nAbstract: 250 words. Main text: 3500 words "
                  "excluding references. References: maximum 50.\n", encoding="utf-8")
    polish("lint")
    r = gate("journal_limits_met")
    record("journal_limits_met parses the quoted caps and passes", r.ok, r.detail[:110])
    gx.write_text("# Guidelines\n\n## Word limits\nAbstract: 250 words. Main text: 100 words "
                  "excluding references. References: maximum 50.\n", encoding="utf-8")
    r = gate("journal_limits_met")
    record("journal_limits_met rejects an over-limit manuscript", not r.ok, r.detail[:110])
    gx.write_text("# Guidelines\n\n## Word limits\nSee the website.\n", encoding="utf-8")
    r = gate("journal_limits_met")
    record("journal_limits_met rejects unquoted limits", not r.ok, r.detail[:110])

    section("polish: allowlist escape hatch")
    disc.write_text(CLEAN + "\nThe guideline states that screening plays a crucial role.\n",
                    encoding="utf-8")
    polish("lint")
    record("allowlist absent -> phrase blocks", not gate("ai_tells_clean").ok)
    (integration / "polish_allowlist.tsv").write_text(
        "plays a crucial role\tquoted verbatim from the 2023 guideline\n", encoding="utf-8")
    polish("lint")
    record("allowlisted phrase stops blocking", gate("ai_tells_clean").ok)

    for f in (disc, gx, integration / "polish_allowlist.tsv"):
        f.unlink(missing_ok=True)
    shutil.rmtree(integration / "prepolish", ignore_errors=True)
    (integration / "polish_report.json").unlink(missing_ok=True)


def run_online(proj: Path) -> None:
    section("live API (network)")
    env = {**os.environ, "MEDPAPER_PROJECT": str(proj), "MEDPAPER_ROOT": str(ROOT),
           "PYTHONIOENCODING": "utf-8"}
    p = subprocess.run([sys.executable, str(ROOT / "tools/pubmed/client.py"),
                        "fetch", "--ids", "32150289", "--with-abstract"],
                       capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("PubMed efetch returns a real record",
           "Sarcopenia Definition" in (p.stdout or ""), (p.stderr or "")[:100])
    p = subprocess.run([sys.executable, str(ROOT / "tools/pubmed/build_library.py"),
                        "--add-ids", "32150289,34315158", "--export"],
                       capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("library builds and exports", (proj / "06_refs/refs.bib").exists(),
           (p.stderr or "")[:100])
    p = subprocess.run([sys.executable, str(ROOT / "tools/pubmed/verify.py")],
                       capture_output=True, text=True, env=env, encoding="utf-8", errors="replace")
    record("verification confirms every entry against the source",
           "verified 2/2" in (p.stdout or ""), (p.stdout or "").strip()[-90:])
    from wfcore import registry
    from wfcore.checks import Ctx, get, load_all
    from wfcore.state import State

    load_all()
    pipe = registry.load()
    state = State(proj, ".wf").load()
    outcome = get("reference_provenance")(Ctx(
        pipeline=pipe,
        state=state,
        project=proj,
        stage=pipe.stage("S13_reflib"),
        spec={"check": "reference_provenance", "live": True},
    ))
    record("independent S13 gate re-fetches verified records live", outcome.ok,
           outcome.detail[:120])


# ---------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description="self test for the medpaper toolchain")
    ap.add_argument("--keep", action="store_true", help="do not delete the temp project")
    ap.add_argument("--online", action="store_true", help="also exercise the live APIs")
    ap.add_argument("--workdir", type=Path,
                    help="use an existing empty directory instead of the system temp root")
    args = ap.parse_args()

    for mod, why in (("matplotlib", "figures"), ("numpy", "QC"), ("openpyxl", "tables"),
                     ("docx", "Word submission files")):
        try:
            __import__(mod)
        except ImportError:
            print(f"cannot run: {mod} is required for {why}.\n"
                  f"  uv pip install --python uv run python {mod}")
            return 1

    if args.workdir:
        tmp = args.workdir.resolve()
        tmp.mkdir(parents=True, exist_ok=True)
        if any(tmp.iterdir()):
            print(f"cannot run: --workdir must be empty: {tmp}")
            return 1
    else:
        tmp = Path(tempfile.mkdtemp(prefix="medpaper_selftest_"))
    proj = tmp / "project"
    os.environ["MEDPAPER_PROJECT"] = str(proj)
    os.environ["MEDPAPER_ROOT"] = str(ROOT)
    os.environ["MEDPAPER_SELFTEST_ROOT"] = str(tmp)
    if args.keep:
        os.environ["MEDPAPER_SELFTEST_KEEP_ARTIFACTS"] = "1"
    print(f"temp project: {proj}\n")

    try:
        section("building fixture")
        build_fixture(proj)
        record("fixture built", True)
        run_checks(proj)
        run_data_proof(proj)
        run_reference_proof(proj)
        run_qc(proj)
        run_archetypes(proj)
        run_codex_integration()
        run_fulltext_registration(proj)
        run_journal_eligibility(proj)
        run_manuscript_docx(proj)
        run_polish(proj)
        if args.online:
            run_online(proj)
    finally:
        if args.keep or args.workdir:
            print(f"\nkept: {tmp}")
        else:
            shutil.rmtree(tmp, ignore_errors=True)

    failed = [r for r in results if r[1] == FAIL]
    print("\n" + "=" * 70)
    print(f"{len(results) - len(failed)}/{len(results)} checks passed")
    if failed:
        for name, _, detail in failed:
            print(f"  FAIL  {name}  {detail}")
        return 2
    print("selftest: OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

